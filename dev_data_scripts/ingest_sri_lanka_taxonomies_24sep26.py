"""Build python/data/sri-lanka-taxonomies.json from the country office's "Sectors for Sri Lanka".

The UNDP country office set Sri Lanka's own lenses in one document (SharePoint, Sri Lanka
folder, Lea Phillips, 24 Sep 2026):

  - Adaptation: nine sectors (Agriculture, Livestock, Fisheries, Biodiversity, Health,
    Tourism, Urban planning and human settlements, Water, Coastal and Marine). They replace
    the GGA areas for Sri Lanka only.
  - Loss and damage: one area, a new lens.
  - Climate mitigation (IPCC sectors), GLOBE and Human rights: "KEEP AS IS", so nothing here.

Names and descriptions are copied, never written. A name is the document's own, the one it
says Sri Lanka uses where it gives one ("Agriculture (called Agriculture and food for Sri
Lanka)" -> "Agriculture and food"), in sentence case. A description is the document's
paragraphs for that row with whitespace collapsed; the "Source:" lines become `_sources`.

Provenance was checked against the public NDC 3.0 PDF (UNFCCC, submitted 22 Sep 2025) and
decision 2/CMA.5 on 6 Oct 2026: the four GGA paragraphs equal categories.json word for word
(the test suite pins this); the paragraphs the document attributes to NDC 3.0 are close
extracts with light edits by the office (US spelling, sentences shortened or joined); the
first loss and damage paragraph (the definition and the attribution caveat) does not appear
in NDC 3.0 and is kept as country-office text (Jonas, 6 Oct 2026).

The document is pinned on its text, not its bytes (SharePoint rewrites the container when
anyone opens it). Dry run by default; --write writes the file.

    python3 dev_data_scripts/ingest_sri_lanka_taxonomies_24sep26.py            # dry run
    python3 dev_data_scripts/ingest_sri_lanka_taxonomies_24sep26.py --write
"""

from __future__ import annotations

import argparse
import glob
import hashlib
import json
import re
from pathlib import Path

import docx  # python-docx

REPO = Path(__file__).resolve().parent.parent
OUT = REPO / "python" / "data" / "sri-lanka-taxonomies.json"

DOC_NAME = "Sectors for Sri Lanka.docx"
DOC_GLOBS = [
    str(Path.home() / "Library/CloudStorage/*AIFlagship*/Sri Lanka" / DOC_NAME),
    str(Path.home() / "Library/CloudStorage/*/*AIFlagship*/Sri Lanka" / DOC_NAME),
]
# sha256 over the document's paragraph and table text (see text_digest), 24 Sep 2026 version.
EXPECTED_DIGEST = "0182179f29902b245c2a16f5c8dbb9495a24f2a906bf815c925d456856fc0826"

GGA_DECISION_URL = "https://unfccc.int/sites/default/files/resource/cma2023_16a01E.pdf#page=9"
NDC_URL = (
    "https://unfccc.int/sites/default/files/2025-09/Sri%20Lankas%20Nationally%20Determined"
    "%20Contributions%203.0%20(2026-2035)%20submitted%2022.09.2025%20(1).pdf"
)

# Document row name -> category id, in the document's order. Pins the rows we expect.
ADAPTATION_IDS = {
    "Agriculture": "lk_agriculture_food",
    "Livestock": "lk_livestock",
    "Fisheries": "lk_fisheries",
    "Biodiversity": "lk_biodiversity",
    "Health": "lk_health",
    "Tourism": "lk_tourism",
    "Urban Planning and Human settlements": "lk_infrastructure_settlements",
    "Water": "lk_water",
    "Coastal and Marine": "lk_coastal_marine",
}
LOSS_DAMAGE_IDS = {"Loss and damage": "lk_loss_damage"}

# The two lenses' attribution, shown as the lens tooltip. English only: Sri Lanka's brief is.
ADAPTATION_TOOLTIP = (
    "Sri Lanka's adaptation sectors, as set by the UNDP country office from the country's "
    "NDC 3.0 (2026-2035) and the Global Goal on Adaptation"
)
LOSS_DAMAGE_TOOLTIP = (
    "Loss and damage, as set by the UNDP country office from Sri Lanka's NDC 3.0 (2026-2035)"
)


def find_doc(explicit: Path | None) -> Path:
    if explicit:
        if not explicit.is_file():
            raise SystemExit(f"--docx not found: {explicit}")
        return explicit
    hits = sorted({h for g in DOC_GLOBS for h in glob.glob(g)})
    if len(hits) != 1:
        raise SystemExit(
            f"could not locate '{DOC_NAME}' under ~/Library/CloudStorage (found {len(hits)}); "
            "pass --docx"
        )
    return Path(hits[0])


def text_digest(d: docx.document.Document) -> str:
    """sha256 over every paragraph and table cell, raw. Stable across SharePoint re-saves."""
    h = hashlib.sha256()
    for p in d.paragraphs:
        h.update((p.text + "\x1e").encode())
    for t in d.tables:
        for row in t.rows:
            h.update(("\x1f".join(c.text for c in row.cells) + "\x1e").encode())
    return h.hexdigest()


def squash(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


def display_name(row_name: str) -> str:
    """The name Sri Lanka uses where the document gives one, in sentence case."""
    called = re.search(r"\(called (.+?) (?:for|in) Sri Lanka\)", row_name)
    name = called.group(1) if called else row_name
    return name[0].upper() + name[1:].lower()


def row_key(row_name: str) -> str:
    return squash(re.sub(r"\(called .+?\)", "", row_name))


def parse_rows(table, ids: dict[str, str]) -> list[dict]:
    cats = []
    for row in table.rows:
        raw_name = squash(row.cells[0].text)
        key = row_key(raw_name)
        if key not in ids:
            raise SystemExit(f"unexpected row {raw_name!r}; expected one of {list(ids)}")
        paragraphs: list[str] = []
        sources: list[str] = []
        pending: list[str] = []
        for line in row.cells[1].text.split("\n"):
            line = squash(line)
            if not line:
                continue
            if line.startswith("Source:"):
                source = squash(line[len("Source:"):])
                sources.extend(f"{source} (paragraph {len(paragraphs) + i + 1})" for i in range(len(pending)))
                paragraphs.extend(pending)
                pending = []
            else:
                pending.append(line)
        if pending:
            raise SystemExit(f"{key}: text after the last 'Source:' line: {pending!r}")
        cats.append(
            {
                "id": ids[key],
                "name": display_name(raw_name),
                "description": "\n\n".join(paragraphs),
                "_documentRow": raw_name,
                "_sources": sources,
            }
        )
    if [c["id"] for c in cats] != list(ids.values()):
        raise SystemExit(f"rows out of the expected order: {[c['id'] for c in cats]}")
    return cats


def build(d: docx.document.Document, digest: str, doc_path: Path) -> dict:
    if len(d.tables) != 2:
        raise SystemExit(f"expected 2 tables (Adaptation, Loss and Damage), found {len(d.tables)}")
    adaptation = parse_rows(d.tables[0], ADAPTATION_IDS)
    loss_damage = parse_rows(d.tables[1], LOSS_DAMAGE_IDS)
    return {
        "_comment": (
            "Sri Lanka's own lenses, set by the UNDP country office. Each entry in `taxonomies` "
            "is classified like a global lens under its own taxonomyType; `replaces` names the "
            "global lens it stands in for on Sri Lanka only. Climate mitigation (IPCC sectors), "
            "Biodiversity (GLOBE) and Human rights stay as they are (the document: 'KEEP AS IS'). "
            "Generated by dev_data_scripts/ingest_sri_lanka_taxonomies_24sep26.py; edit the "
            "document, not this file."
        ),
        "_source": (
            f"'{DOC_NAME}', SharePoint 'GEF GBF Early Action Support OneDrive - AI Flagship/"
            f"Sri Lanka' (Lea Phillips, 24 Sep 2026; text sha256 {digest}). Names are the "
            "document's, using the name it gives for Sri Lanka where it gives one, in sentence "
            "case. Descriptions are the document's paragraphs with whitespace collapsed; each "
            "category's `_sources` carries the document's own source line per paragraph."
        ),
        "taxonomies": [
            {
                "taxonomyType": "adaptation",
                "replaces": "gga",
                "tooltip": ADAPTATION_TOOLTIP,
                "_source": (
                    "Paragraphs attributed to decision 2/CMA.5 equal categories.json "
                    "gga_categories word for word (paragraph 9(a), (b), (c), (e); "
                    f"{GGA_DECISION_URL}). Paragraphs attributed to 'Description in NDC 3.0' "
                    "are close extracts of Sri Lanka's NDC 3.0 with light edits by the country "
                    f"office (checked 6 Oct 2026 against {NDC_URL})."
                ),
                "categories": adaptation,
            },
            {
                "taxonomyType": "loss_damage",
                "tooltip": LOSS_DAMAGE_TOOLTIP,
                "_source": (
                    "Paragraphs 2 and 3 are close extracts of the loss and damage section of "
                    "Sri Lanka's NDC 3.0 (PDF page 79). Paragraph 1, the definition and the "
                    "caveat that a loss is not automatically attributable to climate change, "
                    "does not appear in NDC 3.0 (checked 6 Oct 2026); the document attributes "
                    "it there, and it is kept as country-office text (Jonas Nothnagel, 6 Oct "
                    f"2026). {NDC_URL}"
                ),
                "categories": loss_damage,
            },
        ],
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", type=Path, default=None)
    ap.add_argument("--write", action="store_true")
    args = ap.parse_args()

    path = find_doc(args.docx)
    d = docx.Document(str(path))
    digest = text_digest(d)
    if EXPECTED_DIGEST != "PIN_ME" and digest != EXPECTED_DIGEST:
        raise SystemExit(
            f"'{DOC_NAME}' text changed (sha256 {digest}, pinned {EXPECTED_DIGEST}): "
            "read the new version before re-ingesting"
        )
    out = build(d, digest, path)
    for tax in out["taxonomies"]:
        print(f"{tax['taxonomyType']}: {len(tax['categories'])} categories")
        for c in tax["categories"]:
            print(f"  {c['id']:32} {c['name']!r:42} {len(c['description']):5} chars, {len(c['_sources'])} sources")
    print(f"text sha256 {digest}")
    if args.write:
        OUT.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"wrote {OUT}")
    else:
        print("dry run; --write to write", OUT)


if __name__ == "__main__":
    main()
