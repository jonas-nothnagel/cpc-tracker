"""The review sheet for steps 2 and 3 of the BER plan: a person checks each
line's area, and each line's expenditure type of the official chain.

  build   python/output/<country>/ber_areas_<lens>.json (both models' readings)
          -> one workbook, a sheet per country: every line with its money, each
          model's area and reason, and empty columns for the call. Claude Opus
          decides a line's area, so it is model 1; GPT-5.4 is the second
          opinion. The lines a person should look at first lead (the second
          opinion differs, or no area is clear), largest money first; then the
          rest. Earlier calls (python/data/<country>-ber-areas.json) are
          filled in. Where python/output/<country>/ber_types_globe.json
          exists, a "<Country> types" sheet does the same for the types.
  import  the filled workbook -> python/data/<country>-ber-areas.json (and
          <country>-ber-types.json): each call with who made it and when. Only
          these count as checked; until then the brief marks a line's area or
          type "AI-suggested, not reviewed".

Run from python/:
    .venv/bin/python scripts/ber_review_sheet.py build [--out PATH]
    .venv/bin/python scripts/ber_review_sheet.py import PATH --by "Name"
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.config import DATA_DIR, OUTPUT_DIR  # noqa: E402

COUNTRIES = {"mongolia": "Mongolia", "panama": "Panama", "sri-lanka": "Sri Lanka"}
LENS_KEYS = {"globe": "globe_categories"}
DEFAULT_OUT = Path(__file__).resolve().parent.parent / "tmp" / "ber" / "BER review sheet - GLOBE.xlsx"
CALLS = {"Accept": "accept", "Change": "change", "None fits": "none"}

COLUMNS = [
    ("Line", 16),
    ("Name", 44),
    ("Institution or ministry", 28),
    ("The review's own tags", 22),
    ("Reviewed spending", 16),
    ("US$ million", 12),
    ("Years", 10),
    ("Model 1 (Claude Opus): area", 26),
    ("Model 1: score", 9),
    ("Model 1: reason", 50),
    ("Model 2 (GPT-5.4): area", 26),
    ("Model 2: reason", 50),
    ("The line's area", 22),
    ("Your call", 12),
    ("Area, if you change it", 26),
    ("Note", 30),
]
STATUS_WORDS = {
    "second-differs": "Model 1's area; model 2 differs",
    "no-clear-area": "No specific area",
    "second-agrees": "Model 1's area; model 2 agrees",
    "no-reading": "No reading by model 1",
}
# The lines a person should look at first: model 2 reads them differently, or
# no area reaches the mark.
TO_LOOK_AT = ("second-differs", "no-clear-area")
# Rows where the line takes no area: model 1's area, shown for reference, is
# below the mark (or missing), so an Accept there accepts no specific area.
NO_AREA_WORDS = {STATUS_WORDS["no-clear-area"], STATUS_WORDS["no-reading"]}


def review_tags(line: dict[str, Any]) -> str:
    """The review's own tags on a line: Sri Lanka's list as it is, Panama's
    Directo / Indirecto per year as the tags it carries."""
    tags = line.get("tags")
    if isinstance(tags, list):
        return "; ".join(tags)
    if isinstance(tags, dict):
        return "; ".join(sorted(set(tags.values())))
    return ""


def sheet_order(status: str, money: float, code: str) -> tuple[int, float, str]:
    """Rows a person should look at first, then the rest; each largest money first."""
    rank = 0 if status in TO_LOOK_AT else 1 if status == "no-reading" else 2
    return (rank, -money, code)


def area_names(lens: str) -> dict[str, str]:
    categories = json.loads((DATA_DIR / "categories.json").read_text())[LENS_KEYS[lens]]
    return {c["id"]: c["name"] for c in categories}


def money(record: dict[str, Any], line: dict[str, Any]) -> tuple[float, float, str]:
    """The line's counted amount over the record's complete years, in its own
    currency and in US$ (each year at its own rate)."""
    years = [y for y in record["years"] if y["complete"]]
    local = usd = 0.0
    for y in years:
        cells = (line["amounts"].get(y["year"]) or {}).values()
        value = sum((c or {}).get("counted", 0.0) for c in cells)
        local += value
        scale = 1000.0 if record["unit"] == "billion" else 1.0
        usd += value * scale / y["usdRate"]
    span = f"{years[0]['year']}-{years[-1]['year']}" if years else ""
    return local, usd, span


def build(lens: str, out: Path, countries: tuple[str, ...] = tuple(COUNTRIES), types: bool = True) -> None:
    names = area_names(lens)
    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    for row in [
        ["One row per budget line of the expenditure review, a sheet per country."],
        ["Two AI models placed each line in a GLOBE area from the review's own description of it, without its amounts."],
        ["Model 1 (Claude Opus) decides: a line takes its area at a score of 0.5 or more, else it is general biodiversity spending with no specific area. Model 2 (GPT-5.4) is a second opinion."],
        ["The review's call that a line is biodiversity spending stands in every case: the question here is only which area it serves."],
        ["Rows where model 2 reads the line differently, or no area is clear, come first, largest money first; then the rest."],
        ["Your call: Accept (the line's area as shown: model 1's area, or no specific area where none is clear), Change (pick the area in the next column) or None fits (no specific area)."],
        ["Two models agreeing is not a check: only your call counts as checked. Leave a row empty to decide later."],
        ["Send the sheet back as it is; scripts/ber_review_sheet.py import reads the calls."],
    ]:
        guide.append(row)
    guide.column_dimensions["A"].width = 120

    area_list = list(names.values())
    for country, title in COUNTRIES.items():
        if country not in countries:
            continue
        path = OUTPUT_DIR / country / f"ber_areas_{lens}.json"
        if not path.exists():
            continue
        readings = json.loads(path.read_text())
        record = json.loads((DATA_DIR / f"{country}-ber.json").read_text())["record"]
        by_code = {l["code"]: l for l in record["lines"]}
        earlier = load_calls(country, lens)
        rows = []
        for r in readings["lines"]:
            line = by_code[r["code"]]
            local, usd, span = money(record, line)
            rows.append((r, line, local, usd, span))
        rows.sort(key=lambda x: sheet_order(x[0]["status"], x[2], x[0]["code"]))

        ws = wb.create_sheet(title)
        ws.append([c for c, _ in COLUMNS])
        for i, (_, width) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        unit = f"{record['currency']} {record['unit']}"
        ws.cell(1, 5).value = f"Reviewed spending ({unit})"
        for r, line, local, usd, span in rows:
            # Model 1 is Claude Opus (the second reading), which decides.
            model1, model2 = r["second"] or {}, r["first"] or {}
            call = earlier.get(r["code"], {})
            # An earlier Accept whose area the line no longer takes is filled
            # in as a Change to that area, so the row still says what was
            # decided and a re-import keeps it as it was.
            prefill = call.get("call")
            if prefill == "accept" and call.get("area") != r["area"]:
                prefill = "change" if call.get("area") else "none"
            ws.append(
                [
                    r["code"],
                    line["name"],
                    line.get("institution") or line.get("ministry") or "",
                    review_tags(line),
                    round(local, 3),
                    round(usd, 2),
                    span,
                    names.get(model1.get("area"), "") if model1.get("area") else "",
                    model1.get("score", ""),
                    (model1["ranked"][0]["reasoning"] if model1.get("ranked") else ""),
                    names.get(model2.get("area"), "") if model2.get("area") else "",
                    (model2["ranked"][0]["reasoning"] if model2.get("ranked") else ""),
                    STATUS_WORDS[r["status"]],
                    next((k for k, v in CALLS.items() if v == prefill), ""),
                    names.get(call.get("area"), "") if prefill == "change" else "",
                    call.get("note", ""),
                ]
            )
        shade = PatternFill("solid", fgColor="FFF5C7")
        look_at = {STATUS_WORDS[s] for s in TO_LOOK_AT}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row[12].value in look_at:
                row[12].fill = shade
        last = ws.max_row
        calls = DataValidation(type="list", formula1='"' + ",".join(CALLS) + '"', allow_blank=True)
        areas = DataValidation(type="list", formula1='"' + ",".join(area_list) + '"', allow_blank=True)
        ws.add_data_validation(calls)
        ws.add_data_validation(areas)
        calls.add(f"N2:N{last}")
        areas.add(f"O2:O{last}")
    if types:
        build_type_sheets(wb, countries)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}")


# ---------------------------------------------------------------------------
# The expenditure types of the official chain (step 3)
# ---------------------------------------------------------------------------

TYPE_COLUMNS = [
    ("Line", 16),
    ("Name", 44),
    ("Reviewed spending", 16),
    ("US$ million", 12),
    ("Years", 10),
    ("Model 1: type", 40),
    ("Model 1: score", 9),
    ("Model 1: reason", 50),
    ("Model 2: type", 40),
    ("Model 2: reason", 50),
    ("GBF targets both readings lead to", 18),
    ("The two models", 18),
    ("Your call", 12),
    ("Type code, if you change it", 14),
    ("Note", 30),
]
TYPE_STATUS_WORDS = {
    "models-differ": "pick different types",
    "no-clear-type": "find no clear type",
    "models-agree": "agree",
    "one-model": "one model only",
}


def type_sheet_title(title: str) -> str:
    return f"{title} types"


def build_type_sheets(wb: Workbook, countries: tuple[str, ...] = tuple(COUNTRIES)) -> None:
    """A sheet per country for the lines' expenditure types: rows whose two
    readings lead to no shared GBF target first, then those the models type
    differently, then the rest, largest money first."""
    guide = wb["How to use"]
    guide.append([])
    guide.append(["The 'types' sheets: the same lines, each placed by the two models in a BIOFIN expenditure type (GLOBE level 3)."])
    guide.append(["A type leads to the GBF targets BIOFIN files it under; a line meets a national target of the 7th National Report through a GBF target both readings lead to."])
    guide.append(["Your call: Accept (model 1's type), Change (write the type code, e.g. 7.01.12) or None fits."])
    for country, title in COUNTRIES.items():
        if country not in countries:
            continue
        path = OUTPUT_DIR / country / "ber_types_globe.json"
        if not path.exists():
            continue
        readings = json.loads(path.read_text())
        record = json.loads((DATA_DIR / f"{country}-ber.json").read_text())["record"]
        by_code = {l["code"]: l for l in record["lines"]}
        earlier = load_type_calls(country)
        rows = []
        for r in readings["lines"]:
            local, usd, span = money(record, by_code[r["code"]])
            rows.append((r, by_code[r["code"]], local, usd, span))
        rows.sort(key=lambda x: (x[0]["status"] != "apart", x[0]["typeStatus"] == "models-agree", -x[2], x[0]["code"]))

        ws = wb.create_sheet(type_sheet_title(title))
        ws.append([c for c, _ in TYPE_COLUMNS])
        for i, (_, width) in enumerate(TYPE_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "C2"
        ws.cell(1, 3).value = f"Reviewed spending ({record['currency']} {record['unit']})"

        def named(reading: dict[str, Any] | None) -> str:
            if not reading or not reading.get("ranked"):
                return ""
            top = reading["ranked"][0]
            return f"{top['id']} {top['name']} (GBF {', '.join(str(g) for g in top['gbf'])})"

        for r, line, local, usd, span in rows:
            first, second = r["first"], r["second"]
            call = earlier.get(r["code"], {})
            # As for the areas: an earlier Accept of a type model 1 no longer
            # puts first is filled in as a Change to that type.
            shown_type = first["ranked"][0]["id"] if first["ranked"] else None
            prefill = call.get("call")
            if prefill == "accept" and call.get("type") != shown_type:
                prefill = "change" if call.get("type") else "none"
            ws.append(
                [
                    r["code"],
                    line["name"],
                    round(local, 3),
                    round(usd, 2),
                    span,
                    named(first),
                    first["score"],
                    first["ranked"][0]["reasoning"] if first["ranked"] else "",
                    named(second),
                    second["ranked"][0]["reasoning"] if second and second.get("ranked") else "",
                    ", ".join(str(g) for g in r["gbf"]) or "none",
                    TYPE_STATUS_WORDS[r["typeStatus"]],
                    next((k for k, v in CALLS.items() if v == prefill), ""),
                    (call.get("type") or "") if prefill == "change" else "",
                    call.get("note", ""),
                ]
            )
        shade = PatternFill("solid", fgColor="FFF5C7")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row[10].value == "none" or row[11].value != TYPE_STATUS_WORDS["models-agree"]:
                row[11].fill = shade
        calls = DataValidation(type="list", formula1='"' + ",".join(CALLS) + '"', allow_blank=True)
        ws.add_data_validation(calls)
        calls.add(f"M2:M{ws.max_row}")


def type_calls_path(country: str) -> Path:
    return DATA_DIR / f"{country}-ber-types.json"


def load_type_calls(country: str) -> dict[str, dict[str, Any]]:
    path = type_calls_path(country)
    return json.loads(path.read_text()).get("types", {}) if path.exists() else {}


def line_code(value: Any) -> str:
    """A line's code as the record keys it, whatever the spreadsheet made of
    the cell: a code typed or converted into a number comes back as one."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def kept(earlier: dict[str, Any] | None, entry: dict[str, Any]) -> dict[str, Any]:
    """The call to store: the earlier one, with who made it and when, where
    the sheet repeats its decision (build fills earlier calls in, an Accept
    whose area or type model 1 no longer reads as a Change to it); else the
    new one, by whoever imports it today."""

    def decision(call: dict[str, Any]) -> tuple[Any, Any]:
        value = None if call.get("call") == "none" else call.get("area", call.get("type"))
        return value, call.get("note")

    return earlier if earlier and decision(earlier) == decision(entry) else entry


def import_type_calls(wb: Any, by: str) -> None:
    types = {t["id"] for t in json.loads((DATA_DIR / "globe-expenditure-types.json").read_text())["types"]}
    for country, title in COUNTRIES.items():
        sheet = type_sheet_title(title)
        if sheet not in wb.sheetnames:
            continue
        target = type_calls_path(country)
        stored = json.loads(target.read_text()) if target.exists() else {}
        calls = stored.get("types", {})
        changed = 0
        for row in wb[sheet].iter_rows(min_row=2, values_only=True):
            raw, call, code_given, note = row[0], row[12], row[13], row[14]
            if not raw or not call:
                continue
            code = line_code(raw)
            kind = CALLS.get(str(call).strip())
            if kind is None:
                raise SystemExit(f"{sheet}, line {code}: unknown call {call!r}")
            chosen = str(code_given).strip() if code_given else ""
            if kind == "change" and chosen not in types:
                raise SystemExit(f"{sheet}, line {code}: Change needs a type code from the GLOBE workbook, found {chosen!r}")
            # Accept takes the type the sheet showed the reviewer ("Model 1:
            # type", its code first), not whatever a later run reads.
            shown = str(row[5]).split(" ", 1)[0] if row[5] else ""
            decided = None
            if kind == "change":
                decided = chosen
            elif kind == "accept" and shown in types:
                decided = shown
            entry = {
                "call": kind,
                "type": decided,
                "by": by,
                "at": date.today().isoformat(),
                **({"note": str(note).strip()} if note else {}),
            }
            calls[code] = kept(calls.get(code), entry)
            if calls[code] is entry:
                changed += 1
        stored["_provenance"] = (
            "A person's calls on the expenditure types the AI suggested for the review's lines (step 3 of the BER "
            "plan, the official chain); written by python/scripts/ber_review_sheet.py import. Only these count as checked."
        )
        stored["types"] = dict(sorted(calls.items()))
        target.write_text(json.dumps(stored, indent=2, ensure_ascii=False) + "\n")
        print(f"{sheet}: {changed} calls new or changed, {len(calls)} in all -> {target}")


def calls_path(country: str) -> Path:
    return DATA_DIR / f"{country}-ber-areas.json"


def load_calls(country: str, lens: str) -> dict[str, dict[str, Any]]:
    path = calls_path(country)
    if not path.exists():
        return {}
    return json.loads(path.read_text()).get(lens, {})


def import_calls(path: Path, lens: str, by: str) -> None:
    names = area_names(lens)
    by_name = {v: k for k, v in names.items()}
    wb = load_workbook(path, data_only=True)
    for country, title in COUNTRIES.items():
        if title not in wb.sheetnames:
            continue
        ws = wb[title]
        target = calls_path(country)
        stored = json.loads(target.read_text()) if target.exists() else {}
        calls = stored.get(lens, {})
        changed = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw, call, area, note = row[0], row[13], row[14], row[15]
            if not raw or not call:
                continue
            code = line_code(raw)
            kind = CALLS.get(str(call).strip())
            if kind is None:
                raise SystemExit(f"{title}, line {code}: unknown call {call!r}")
            if kind == "change" and area not in by_name:
                raise SystemExit(f"{title}, line {code}: Change needs an area from the list, found {area!r}")
            # Accept takes the area the sheet showed the reviewer ("Model 1
            # (Claude Opus): area"), not whatever a later run of the models reads;
            # on a row with no specific area it accepts just that.
            decided = None
            if kind == "change":
                decided = by_name[area]
            elif kind == "accept" and row[12] not in NO_AREA_WORDS:
                decided = by_name.get(row[7])
            entry = {
                "call": kind,
                "area": decided,
                "by": by,
                "at": date.today().isoformat(),
                **({"note": str(note).strip()} if note else {}),
            }
            calls[code] = kept(calls.get(code), entry)
            if calls[code] is entry:
                changed += 1
        stored["_provenance"] = (
            "A person's calls on the areas the AI suggested for the review's lines (step 2 of the BER plan); "
            "written by python/scripts/ber_review_sheet.py import. Only these count as checked."
        )
        stored[lens] = dict(sorted(calls.items()))
        target.write_text(json.dumps(stored, indent=2, ensure_ascii=False) + "\n")
        print(f"{title}: {changed} calls new or changed, {len(calls)} in all -> {target}")
    import_type_calls(wb, by)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    b = sub.add_parser("build")
    b.add_argument("--lens", default="globe", choices=sorted(LENS_KEYS))
    b.add_argument("--out", type=Path, default=DEFAULT_OUT)
    b.add_argument("--country", choices=[*COUNTRIES, "all"], default="all", help="One country's sheet only, e.g. to send to its review team")
    b.add_argument("--areas-only", action="store_true", help="Leave out the expenditure-type sheets")
    i = sub.add_parser("import")
    i.add_argument("sheet", type=Path)
    i.add_argument("--lens", default="globe", choices=sorted(LENS_KEYS))
    i.add_argument("--by", required=True, help="Who made the calls, as the brief should name them")
    args = parser.parse_args()
    if args.command == "build":
        build(args.lens, args.out, tuple(COUNTRIES) if args.country == "all" else (args.country,), types=not args.areas_only)
    else:
        import_calls(args.sheet, args.lens, args.by)


if __name__ == "__main__":
    main()
