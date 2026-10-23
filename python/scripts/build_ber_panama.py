"""Build Panama's spending record from its Biodiversity Expenditure Review (BER).

Step 1 of the BER plan, "Show the money". Reads BIOFIN Panama's budget
database (Base de datos FINAL.xlsx: every budget line of the 34 screened
institutions, 2015-2024, with BIOFIN's tag) and BIOFIN's published tables
(Tablas_adicionales_desagregacion_descripciones (1).xlsx), and writes
python/data/panama-ber.json. No AI touches a figure, a name or a description.

The review's rule, as BIOFIN's own pivot in the database computes it (sheet
Hoja1, filtered to its tags): reviewed spending = the committed amount
(presupuesto_ejecutado, the Ministry of Economy and Finance's "compromiso") of
every line tagged Directo or Indirecto, with no weighting. Lines tagged
Saneamiento are the review's special case, kept apart, never added.

Every year, and every institution in every year, must equal BIOFIN's pivot to
the cent, and every year its published headline, or nothing is written.

The record's lines are the budget activities (activity or project), the level
at which BIOFIN tagged them; each sits in its programme and institution. The
series shown starts in 2016: the Ministry of Environment, created in 2015 out
of ANAM, has no lines in 2015.

Run from python/ (a few minutes the first time; the extract is cached in
python/tmp/ber/):
    .venv/bin/python scripts/build_ber_panama.py
"""

from __future__ import annotations

import argparse
import glob
import os
import pickle
import re
import sys
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.ber_record import (  # noqa: E402
    RECORD_VERSION,
    Check,
    amount,
    checks_json,
    berdata_view,
    reconciliation_markdown,
    require_reproduced,
    write_record,
)
from src.config import DATA_DIR  # noqa: E402
from src.xlsx_stream import Workbook  # noqa: E402

FOLDER = "Panama/Panama Materials/Financial data"
DATABASE = "Base de datos FINAL.xlsx"
TABLES = "Tablas_adicionales_desagregacion_descripciones (1).xlsx"
OUT = DATA_DIR / "panama-ber.json"
CACHE = Path(__file__).resolve().parent.parent / "tmp" / "ber"

CHECKED_YEARS = [str(y) for y in range(2015, 2025)]
SHOWN_YEARS = [str(y) for y in range(2016, 2025)]
REVIEWED_TAGS = ("Directo", "Indirecto")
SPECIAL_TAG = "Saneamiento"
M = 1_000_000.0
CENT = 0.005  # half a cent, in balboas

# The database names two institutions differently from BIOFIN's tables: the
# agricultural research institute was renamed in 2021, and one name carries a
# typo. These three are the only names not matched letter for letter.
TABLE_NAME_OF = {
    "INST. DE INVESTIGACIONES AGROPECUARIAS": "Instituto de Investigaciones Agropecuarias",
    "INSTITUTO DE INNOVACIÓN AGROPECUARIA": "Instituto de Investigación Agropecuaria",
    "MINSTERIO DE CULTURA": "Ministerio de Cultura",
}

# Columns of sheet BD (header row 1).
COLS = {
    "año": "A", "area": "B", "desc_area": "C", "entidad": "D", "desc_ent": "E", "tipo": "F", "desc_tipo": "G",
    "programa": "H", "desc_prog": "I", "subprograma": "J", "desc_subprog": "K", "actividad": "L", "desc_act": "M",
    "fuente": "N", "desc_fuente": "O", "desc_grupo": "Q", "desc_economica": "W", "desc_grupo_sec": "AG",
    "desc_cta_sec": "AI", "desc_sectorial": "AK", "Clasificacion Biofin": "AL", "ley": "AM",
    "presupuesto_modificado": "AO", "presupuesto_ejecutado": "AP", "presupuesto_pagado": "AQ", "devengado": "AR",
    "Ejecutado_ponderado": "AV",
}

FUNDING_GROUPS = {
    # The Ministry of Economy and Finance's funding-source groups (source code / 100).
    0: "Current revenue", 1: "Current revenue", 2: "Patrimony", 3: "External loans", 4: "Grants",
    5: "Capital transfers", 6: "Domestic credit", 7: "Canal tolls and dividends",
}


def norm(text: str) -> str:
    """Case, accents and spacing set aside, for matching one BIOFIN table's names to another's."""
    flat = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return " ".join(flat.upper().split())


def text(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def money(value: Any) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


# Module-level factories, so a cached extract can be pickled.
def _floats() -> defaultdict:
    return defaultdict(float)


def _nested() -> defaultdict:
    return defaultdict(_floats)


def _counts() -> defaultdict:
    return defaultdict(int)


@dataclass
class Activity:
    key: tuple
    names: dict[str, dict[str, str]] = field(default_factory=dict)  # year -> names at each level
    years: dict[str, dict[str, float]] = field(default_factory=_nested)
    tags: dict[str, str] = field(default_factory=dict)  # year -> Directo | Indirecto
    sector: dict[tuple[str, str, str], float] = field(default_factory=_floats)
    funding: dict[str, float] = field(default_factory=_floats)
    objects: dict[str, float] = field(default_factory=_floats)


@dataclass
class Extract:
    activities: dict[tuple, Activity] = field(default_factory=dict)
    screened: dict[str, float] = field(default_factory=_floats)  # year -> all lines, committed
    screened_by_entity: dict[tuple[str, str], float] = field(default_factory=_floats)
    special: dict[str, dict[str, float]] = field(default_factory=_nested)
    reviewed_by_entity: dict[tuple[str, str], dict[str, float]] = field(default_factory=_nested)
    tags_seen: dict[str, int] = field(default_factory=_counts)


def read_database(wb: Workbook) -> Extract:
    out = Extract()
    header = None
    for row, v in wb.rows("BD"):
        if row == 1:
            header = {name: v.get(col) for name, col in COLS.items()}
            wrong = {n: h for n, h in header.items() if h != n}
            if wrong:
                raise SystemExit(f"Sheet BD's header has changed: {wrong}")
            continue
        year = str(v.get("A"))
        if year not in CHECKED_YEARS:
            continue
        tag = text(v.get("AL"))
        out.tags_seen[tag] += 1
        entity = text(v.get("E"))
        committed = money(v.get("AP"))
        out.screened[year] += committed
        out.screened_by_entity[(entity, year)] += committed
        if tag == SPECIAL_TAG:
            out.special[year]["committed"] += committed
            out.special[year]["weighted"] += money(v.get("AV"))
            continue
        if tag not in REVIEWED_TAGS:
            continue
        by_entity = out.reviewed_by_entity[(entity, year)]
        by_entity["committed"] += committed
        by_entity["modified"] += money(v.get("AO"))
        # An activity is its codes, never its name: names change across years
        # and the same name recurs in different programmes.
        key = (v.get("B"), v.get("D"), v.get("F"), v.get("H"), v.get("J"), v.get("L"))
        act = out.activities.get(key)
        if act is None:
            act = out.activities[key] = Activity(key)
        act.names[year] = {
            "entity": entity,
            "type": text(v.get("G")),
            "programme": text(v.get("I")),
            "subprogramme": text(v.get("K")),
            "activity": text(v.get("M")),
            "area": text(v.get("C")),
        }
        cell = act.years[year]
        cell["law"] += money(v.get("AM"))
        cell["modified"] += money(v.get("AO"))
        cell["committed"] += committed
        cell["accrued"] += money(v.get("AR"))
        cell["paid"] += money(v.get("AQ"))
        previous = act.tags.get(year)
        if previous and previous != tag:
            raise SystemExit(f"Activity {key} carries both {previous} and {tag} in {year}")
        act.tags[year] = tag
        act.sector[(text(v.get("AG")), text(v.get("AI")), text(v.get("AK")))] += committed
        source = v.get("N")
        group = FUNDING_GROUPS.get(int(source) // 100, "Other") if isinstance(source, (int, float)) else "Other"
        act.funding[group] += committed
        act.objects[text(v.get("Q"))] += committed
    if header is None:
        raise SystemExit("Sheet BD has no header row")
    return out


def read_pivot(wb: Workbook) -> tuple[dict[tuple[str, str], float], dict[str, float]]:
    """BIOFIN's own pivot (sheet Hoja1): committed amount by institution and year."""
    year_col: dict[str, str] = {}
    by_entity: dict[tuple[str, str], float] = {}
    totals: dict[str, float] = {}
    filter_ok = False
    for row, v in wb.rows("Hoja1"):
        if text(v.get("A")) == "Clasificacion Biofin":
            filter_ok = text(v.get("B")) == "(Multiple Items)"
        if not year_col and any(isinstance(v.get(c), int) and 2015 <= v.get(c) <= 2024 for c in v):
            year_col = {c: str(y) for c, y in v.items() if isinstance(y, int) and 2015 <= y <= 2024}
            continue
        label = text(v.get("A"))
        if not year_col or not label or label == "Row Labels":
            continue
        for col, year in year_col.items():
            value = money(v.get(col))
            if label == "Grand Total":
                totals[year] = value
            else:
                by_entity[(label, year)] = value
    if not filter_ok or not totals:
        raise SystemExit("Hoja1 is not BIOFIN's pivot filtered to its tags")
    return by_entity, totals


def read_tables(path: Path) -> tuple[dict[str, dict[str, float]], dict[str, dict[str, tuple[float, float]]]]:
    """BIOFIN's published tables: the headline with the public sector's
    spending and GDP by year, and each institution's modified and committed
    amounts by year (millions of balboas). The tables' "Total general" row
    is left out: the year checks already hold the totals."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    headline: dict[str, dict[str, float]] = {}
    for row in wb["Gasto total en BD"].iter_rows(values_only=True):
        if isinstance(row[0], int) and 2015 <= row[0] <= 2024:
            headline[str(row[0])] = {"reviewed": float(row[1]), "spnf": float(row[2]), "gdp": float(row[4])}
    institutions: dict[str, dict[str, tuple[float, float]]] = {}
    rows = list(wb["Presupuesto asignado a BD"].iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[1] == "Presupuesto Modificado (vigente)")
    years_row = rows[start - 1]
    pairs = [(i, str(y)) for i, y in enumerate(years_row) if isinstance(y, int)]
    for r in rows[start + 1 :]:
        if not r[0] or str(r[0]).strip().lower().startswith("total"):
            break
        institutions[str(r[0]).strip()] = {year: (float(r[i] or 0), float(r[i + 1] or 0)) for i, year in pairs}
    return headline, institutions


def latest(names: dict[str, dict[str, str]], level: str) -> tuple[str, list[str]]:
    """The name in the latest year, and the other names the codes carried."""
    ordered = [names[y][level] for y in sorted(names)]
    current = ordered[-1]
    others = []
    for n in reversed(ordered):
        if n != current and n not in others:
            others.append(n)
    return current, others


def build(database: Path, tables: Path) -> dict[str, Any]:
    stat = database.stat()
    CACHE.mkdir(parents=True, exist_ok=True)
    cached = CACHE / f"panama-extract-{stat.st_size}-{int(stat.st_mtime)}-v{RECORD_VERSION}.pkl"
    if cached.exists():
        print(f"Database from the cached extract {cached.name}")
        ex, pivot_by_entity, pivot_totals = pickle.loads(cached.read_bytes())
    else:
        with Workbook(database) as wb:
            ex = read_database(wb)
            pivot_by_entity, pivot_totals = read_pivot(wb)
        cached.write_bytes(pickle.dumps((ex, pivot_by_entity, pivot_totals)))
    headline, institutions = read_tables(tables)
    print(f"Tags in sheet BD: {dict(ex.tags_seen)}")

    # Institution names as BIOFIN publishes them.
    table_by_norm = {norm(name): name for name in institutions}
    display: dict[str, str] = {}
    for entity, _ in list(ex.reviewed_by_entity) + list(ex.screened_by_entity):
        if entity in display:
            continue
        published = TABLE_NAME_OF.get(entity) or table_by_norm.get(norm(entity))
        display[entity] = published or entity.title()

    checks: list[Check] = []
    db = DATABASE
    for year in CHECKED_YEARS:
        reviewed = sum(a.years[year]["committed"] for a in ex.activities.values() if year in a.years)
        checks.append(Check(f"{year}, all reviewed lines", f"{db}, sheet Hoja1 (BIOFIN's pivot), Grand Total", pivot_totals[year] / M, reviewed / M, CENT / M))
        checks.append(Check(f"{year}, headline", f"{TABLES}, sheet Gasto total en BD", headline[year]["reviewed"], reviewed / M, CENT / M))
    for (entity, year), value in sorted(pivot_by_entity.items()):
        got = ex.reviewed_by_entity.get((entity, year), {}).get("committed", 0.0)
        checks.append(Check(f"{year}, {display.get(entity, entity)}", f"{db}, sheet Hoja1", value / M, got / M, CENT / M, rule="every institution, every year, against BIOFIN's pivot (sheet Hoja1)"))
    for name, years in institutions.items():
        entities = [e for e, d in display.items() if d == name]
        for year, (modified, committed) in years.items():
            got_c = sum(ex.reviewed_by_entity.get((e, year), {}).get("committed", 0.0) for e in entities)
            got_m = sum(ex.reviewed_by_entity.get((e, year), {}).get("modified", 0.0) for e in entities)
            checks.append(Check(f"{year}, {name}, committed", f"{TABLES}, sheet Presupuesto asignado a BD", committed, got_c / M, CENT / M, rule="every institution's committed amount, every year, in BIOFIN's published table"))
            checks.append(Check(f"{year}, {name}, modified budget", f"{TABLES}, sheet Presupuesto asignado a BD", modified, got_m / M, CENT / M, rule="every institution's modified budget, every year, in BIOFIN's published table"))
    require_reproduced(checks)

    # Lines: the activities with money in the years shown.
    programmes: dict[tuple, dict[str, Any]] = {}
    lines = []
    for act in ex.activities.values():
        shown = {y: c for y, c in act.years.items() if y in SHOWN_YEARS}
        if not any(abs(c["committed"]) > 0 for c in shown.values()):
            continue
        area, entity_code, type_code, prog, sub, actv = act.key
        names_years = {y: n for y, n in act.names.items() if y in SHOWN_YEARS} or act.names
        name, aliases = latest(names_years, "activity")
        entity, _ = latest(names_years, "entity")
        programme_name, programme_aliases = latest(names_years, "programme")
        subprogramme, _ = latest(names_years, "subprogramme")
        budget_type, _ = latest(names_years, "type")
        programme_code = f"{area}-{entity_code}-{type_code}-{prog}"
        code = f"{programme_code}-{sub}-{actv}"
        sector = max(act.sector.items(), key=lambda kv: kv[1])[0] if act.sector else ("", "", "")
        total = sum(c["committed"] for c in shown.values())
        programmes.setdefault(
            (area, entity_code, type_code, prog),
            {"code": programme_code, "name": programme_name, "type": budget_type, "institution": display[entity]},
        )
        lines.append(
            {
                "code": code,
                "kind": "activity",
                "name": name,
                **({"formerNames": aliases} if aliases else {}),
                "description": "",
                "source": f"{DATABASE}, sheet BD (desc_act)",
                "institution": display[entity],
                "programme": programme_code,
                "programmeName": programme_name,
                "subprogramme": subprogramme,
                "budgetType": budget_type,
                "sector": {"group": sector[0], "account": sector[1], "sector": sector[2]},
                "tags": {y: t for y, t in sorted(act.tags.items()) if y in SHOWN_YEARS},
                "amounts": {
                    y: {
                        "all": {
                            "plan": amount(c["modified"] / M),
                            "actual": amount(c["committed"] / M),
                            "counted": amount(c["committed"] / M),
                            "accrued": amount(c["accrued"] / M),
                        }
                    }
                    for y, c in sorted(shown.items())
                },
                "funding": [
                    {"source": s, "share": round(v / total, 4)}
                    for s, v in sorted(act.funding.items(), key=lambda kv: -kv[1])
                    if total > 0 and v > 0
                ],
                "payments": [
                    {"name": s.title(), "share": round(v / total, 4)}
                    for s, v in sorted(act.objects.items(), key=lambda kv: -kv[1])[:6]
                    if total > 0 and v > 0
                ],
            }
        )
    lines.sort(key=lambda l: (-sum(c["all"]["counted"] for c in l["amounts"].values()), l["code"]))

    years_out = []
    for year in SHOWN_YEARS:
        reviewed = sum(a.years[year]["committed"] for a in ex.activities.values() if year in a.years)
        years_out.append(
            {
                "year": year,
                "levels": ["all"],
                "complete": True,
                "reviewed": {"all": amount(reviewed / M)},
                "public": {"all": amount(headline[year]["spnf"])},
                "screened": amount(ex.screened[year] / M),
                "gdp": amount(headline[year]["gdp"]),
                "usdRate": 1.0,
                "special": {
                    "committed": amount(ex.special[year]["committed"] / M),
                    "weighted": amount(ex.special[year]["weighted"] / M),
                },
            }
        )
    reviewed_2015 = sum(a.years["2015"]["committed"] for a in ex.activities.values() if "2015" in a.years)
    programs, expenditure = berdata_view(lines, SHOWN_YEARS)
    for p, line in zip(programs, lines):
        p["nameEs"] = line["name"]
        p["institution"] = line["institution"]
    for e, line in zip(expenditure, lines):
        e["institution"] = line["institution"]
    record = {
        "version": RECORD_VERSION,
        "review": {
            "name": "Biodiversity Expenditure Review (BER)",
            "title": "Revisión del gasto público en biodiversidad, Panamá 2015-2024",
            "preparedBy": "BIOFIN Panama (UNDP)",
            "date": "2026-04",
            "files": [DATABASE, TABLES],
            "workbookYears": [CHECKED_YEARS[0], CHECKED_YEARS[-1]],
        },
        "currency": "PAB",
        "unit": "million",
        "usd": {"source": "The balboa is pegged to the US$ at par."},
        "levels": [{"id": "all", "name": "Screened institutions"}],
        "years": years_out,
        "programmes": sorted(programmes.values(), key=lambda p: p["code"]),
        "lines": lines,
        "checks": checks_json(checks),
        "notes": [
            f"Reviewed spending is the committed amount (compromiso) of every budget line BIOFIN tagged Directo or Indirecto, counted in full. 2015-2024: PAB {sum(headline[y]['reviewed'] for y in CHECKED_YEARS):.2f} million.",
            f"2015 is left out of the series: the Ministry of Environment, created that year out of ANAM, has no lines in 2015, and from 2016 it holds 42 to 64% of each year's reviewed spending. 2015 holds PAB {reviewed_2015 / M:.2f} million.",
            "Lines tagged Saneamiento (environmental health and sanitation) are the review's special case, counted at 50% outside the main series; they are kept apart here and never added.",
            "All public spending is the non-financial public sector's accrued spending (SPNF), as BIOFIN's tables give it; reviewed spending is committed, so the two stages differ by about 4%.",
            "Lines are budget activities, the level at which BIOFIN tagged them, each in its programme and institution; activity names are the budget database's own, in Spanish.",
            "Each line keeps the database's sectoral classification as a label. Its code Administración General covers 30 activities (PAB 211 million), forest restoration and water resources among them, so it does not mark administration and is not used to.",
        ],
    }
    return {
        "_provenance": (
            f"Built by python/scripts/build_ber_panama.py from BIOFIN Panama's budget database ({DATABASE}) and published tables ({TABLES}). "
            "Amounts follow the review's own rule; every published total is reproduced to the cent (record.checks). "
            "No AI touched any figure, name or description."
        ),
        "programs": programs,
        "expenditure": expenditure,
        "currency": "PAB",
        "unit": "million",
        "period": {"start": int(SHOWN_YEARS[0]), "end": int(SHOWN_YEARS[-1])},
        "record": record,
    }


def sharepoint_file(name: str) -> Path:
    roots = glob.glob(os.path.expanduser("~/Library/CloudStorage/*/*AIFlagship*"))
    if not roots:
        raise SystemExit("The SharePoint folder is not mounted; pass --database and --tables")
    return Path(roots[0]) / FOLDER / name


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--database", type=Path, help=f"{DATABASE} (default: the SharePoint copy)")
    parser.add_argument("--tables", type=Path, help=f"{TABLES} (default: the SharePoint copy)")
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--check-only", action="store_true", help="Reconcile and report; write nothing")
    args = parser.parse_args()
    payload = build(args.database or sharepoint_file(DATABASE), args.tables or sharepoint_file(TABLES))
    record = payload["record"]
    ok = sum(1 for c in record["checks"] if c["ok"])
    print(f"{ok} of {len(record['checks'])} published totals reproduced; {len(record['lines'])} activities in {len(record['programmes'])} programmes")
    if not args.check_only:
        write_record(args.out, payload)
        (CACHE / "panama-reconciliation.md").write_text(reconciliation_markdown("Panama", record))
        print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
