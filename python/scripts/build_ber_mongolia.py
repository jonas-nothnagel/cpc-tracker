"""Build Mongolia's spending record from its Biodiversity Expenditure Review (BER).

Step 1 of the BER plan, "Show the money". Reads the review's own workbook
(BER_1112.xlsx, Mongolian, with formulas) and its English report (BER [ENG]
v1.2.docx, Unet Consulting for UNDP BIOFIN, November 2025) and writes
python/data/mongolia-ber.json. No AI touches a figure, a name or a description.

The review's rule, as its workbook computes it (sheet Хөтөлбөр, one SUMIFS per
line and year; the adapter checks the formulas say so before it reads a row):

  counted = BAR rate of the line item's purpose code x actual
  - the 13 environment programmes: every line item of programme 714xx,
    whatever its purpose;
  - 15 environmental purpose codes outside them (817xx, 82212), at 100%.

The workbook's sheets hold different parts of the budget: 2016-2021 the State
(central) and local budgets, 2022 and 2024 the local budgets only, 2023 the
State budget only. The record keeps the two levels apart and gives a year
total only where both are in the data (2020 and 2021).

Every line, every year 2016-2024, must equal the workbook's own figure on
sheet Хөтөлбөр, and every year total its row 29 (report Table 9), or nothing
is written.

Each line also keeps who holds its budget in 2020 and 2021 (the 2016-2021
sheet's Admin1 column, in its own English and Mongolian: ministers, the
capital city, the provinces); in each year the holders' amounts must add up
to the line's counted amount, or nothing is written.

Run from python/ (about five minutes the first time; the extract is cached in
python/tmp/ber/):
    .venv/bin/python scripts/build_ber_mongolia.py
"""

from __future__ import annotations

import argparse
import glob
import os
import pickle
import re
import sys
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.ber_record import (  # noqa: E402
    RECORD_VERSION,
    Check,
    amount,
    berdata_view,
    checks_json,
    reconciliation_markdown,
    require_reproduced,
    usd_rates,
    write_record,
)
from src.config import DATA_DIR  # noqa: E402
from src.xlsx_stream import Workbook  # noqa: E402

FOLDER = "Mongolia/Mongolia materials/Finance information/Biodiversity finance"
WORKBOOK = "BER_1112.xlsx"
REPORT = "BER [ENG] v1.2.docx"
OUT = DATA_DIR / "mongolia-ber.json"
CACHE = Path(__file__).resolve().parent.parent / "tmp" / "ber"

PROGRAM_SHEET = "Хөтөлбөр"
PURPOSE_SHEET = "Зориулалт"
GDP_SHEET = "GDP"

CHECKED_YEARS = [str(y) for y in range(2016, 2025)]
SHOWN_YEARS = [str(y) for y in range(2020, 2025)]
LEVELS = [("central", "State budget"), ("local", "Local budgets")]
# The part of the budget each sheet holds (the 2016 sheet's local line items
# carry no programme or purpose code, so 2016 is central in effect).
SHEET_LEVELS = {
    "2016": ["central"],
    **{str(y): ["central", "local"] for y in range(2017, 2022)},
    "2022": ["local"],
    "2023": ["central"],
    "2024": ["local"],
}
COMPLETE_YEARS = [y for y in SHOWN_YEARS if SHEET_LEVELS[y] == ["central", "local"]]

# Bodies kept per line and year for the drill-down; the rest are counted.
BODIES_KEPT = 15
PAYMENTS_KEPT = 6
# The extract's own version, in its cache name: 2 adds the budget holders.
EXTRACT_VERSION = 2

BN = 1_000_000_000.0


# ---------------------------------------------------------------------------
# Excel's SUMIFS criteria, as the workbook applies them
# ---------------------------------------------------------------------------


def excel_text(value: Any) -> str:
    """A cell as SUMIFS compares it with a text criterion."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip() if value is not None else ""


def equals(cell: Any, criterion: str) -> bool:
    """SUMIFS equality: case-insensitive text; numbers match their digits."""
    return excel_text(cell).casefold() == criterion.casefold()


def not_714(cell: Any) -> bool:
    """SUMIFS "<>714*": a number never matches a text wildcard, so it passes."""
    return not (isinstance(cell, str) and cell.startswith("714"))


def number(value: Any) -> float:
    """SUMIFS sums numbers only: text such as "-" counts as nothing."""
    return float(value) if isinstance(value, (int, float)) else 0.0


def code_of(label: Any) -> str:
    """The code that opens a budget label: "71405 Тусгай ..." -> "71405"."""
    text = excel_text(label)
    return text.split()[0] if text else ""


# ---------------------------------------------------------------------------
# The rule sheets: which lines, which rates, which published figures
# ---------------------------------------------------------------------------


@dataclass
class Rule:
    """What the workbook says before any line item is read."""

    programmes: list[str] = field(default_factory=list)  # 714xx, sheet order
    programme_labels: dict[str, str] = field(default_factory=dict)  # code -> Mongolian label (the T criterion)
    purposes: list[str] = field(default_factory=list)  # the codes outside programmes, first occurrence
    rates: dict[str, float] = field(default_factory=dict)  # purpose -> BAR, first match like XLOOKUP
    purpose_names_mn: dict[str, str] = field(default_factory=dict)
    # The workbook's own Mongolian descriptions: 8 of the 13 programmes (sheet
    # Хөтөлбөр, column M) and the purpose codes' definitions (Зориулалт, column O).
    descriptions_mn: dict[str, str] = field(default_factory=dict)
    published: dict[tuple[str, str], float] = field(default_factory=dict)  # (line, year) -> bn
    published_total: dict[str, float] = field(default_factory=dict)  # year -> bn (row 29)
    gdp: dict[str, float] = field(default_factory=dict)  # year -> bn
    formulas: dict[str, str] = field(default_factory=dict)


EXPECTED_FORMULAS = {
    # cell: fragments the review's formula must contain (whitespace removed)
    "C14": ["SUMIFS('2016-2021'!$AB:$AB", "'2016-2021'!$T:$T,$B14", "'2016-2021'!$A:$A,C$13"],
    "I14": ["SUMIFS('2022'!X:X", "'2022'!T:T,Хөтөлбөр!A14"],
    "J14": ["SUMIFS('2023'!I:I", "'2023'!B:B,A14"],
    "K14": ["SUMIFS('2024'!AF:AF", "'2024'!Y:Y,Хөтөлбөр!A14"],
    "C32": ["SUMIFS('2016-2021'!$AB:$AB", "'2016-2021'!$S:$S,\"<>714*\"", "'2016-2021'!$X:$X,Хөтөлбөр!$A32"],
    "I32": ["SUMIFS('2022'!X:X", "'2022'!V:V,Хөтөлбөр!A32", "'2022'!T:T,\"<>714*\""],
    "J32": ["SUMIFS('2023'!I:I", "'2023'!B:B,Хөтөлбөр!$A32", "'2023'!A:A,\"<>714*\""],
    "K32": ["SUMIFS('2024'!AF:AF,'2024'!Z:Z,A32)", "-SUMIFS('2024'!AF:AF,'2024'!Z:Z,A32,'2024'!Y:Y,\">=71400\",'2024'!Y:Y,\"<=71499\")"],
    "C29": ["SUM(C27:C28)"],
}


def read_rule(wb: Workbook) -> Rule:
    rule = Rule()
    formulas = {ref: f for ref, f in wb.formulas(PROGRAM_SHEET) if ref in EXPECTED_FORMULAS}
    for ref, fragments in EXPECTED_FORMULAS.items():
        text = re.sub(r"\s+", "", formulas.get(ref, ""))
        missing = [frag for frag in fragments if re.sub(r"\s+", "", frag) not in text]
        if missing:
            raise SystemExit(
                f"{PROGRAM_SHEET}!{ref} no longer holds the rule this adapter mirrors.\n"
                f"  formula: {formulas.get(ref)!r}\n  missing: {missing}"
            )
        rule.formulas[ref] = formulas[ref]

    year_col: dict[str, str] = {}
    section = "programmes"
    for row, values in wb.rows(PROGRAM_SHEET):
        if not year_col:
            # Rows above the year header describe the programmes in Mongolian.
            if values.get("C") == 2016:
                year_col = {col: str(values[col]) for col in "CDEFGHIJK" if isinstance(values.get(col), int)}
            elif re.fullmatch(r"714\d\d", excel_text(values.get("A"))) and excel_text(values.get("M")):
                rule.descriptions_mn[excel_text(values.get("A"))] = " ".join(excel_text(values.get("M")).split())
            continue
        label = excel_text(values.get("B"))
        code = excel_text(values.get("A"))
        if label == "Нийт байгаль орчны зардал":
            rule.published_total = {year: float(values[col]) for col, year in year_col.items()}
            continue
        if code == "Хүрээлэн буй орчны бус":
            section = "purposes"
            continue
        if not re.fullmatch(r"\d{5}", code):
            continue
        published = {year: float(values.get(col, 0) or 0) for col, year in year_col.items()}
        if section == "programmes" and code.startswith("714"):
            rule.programmes.append(code)
            rule.programme_labels[code] = label
        elif section == "purposes":
            if code in rule.purposes:
                continue  # listed twice on the sheet (81709, 81702, 81705); the sum row counts zero twice
            rule.purposes.append(code)
        else:
            continue
        for year, value in published.items():
            rule.published[(code, year)] = value

    for row, values in wb.rows(PURPOSE_SHEET):
        code = code_of(values.get("A"))
        if not re.fullmatch(r"\d{5}", code):
            continue
        rate = values.get("B")
        if code not in rule.rates and isinstance(rate, (int, float)):
            rule.rates[code] = float(rate)  # XLOOKUP returns the first match
        name = excel_text(values.get("M"))
        name = re.sub(r"^\d{5,6}\s+", "", name)
        if name and code not in rule.purpose_names_mn:
            rule.purpose_names_mn[code] = name
        definition = " ".join(excel_text(values.get("O")).split())
        # Column O holds a definition in the upper block; the lower block repeats the rate there.
        if definition and not re.fullmatch(r"[\d.]+", definition) and code not in rule.descriptions_mn:
            rule.descriptions_mn[code] = definition

    for row, values in wb.rows(GDP_SHEET):
        if excel_text(values.get("C")) == "Нэрлэсэн ДНБ":
            header = {"D": "2016", "E": "2017", "F": "2018", "G": "2019", "H": "2020", "I": "2021", "J": "2022", "K": "2023", "L": "2024"}
            rule.gdp = {year: float(values[col]) for col, year in header.items() if isinstance(values.get(col), (int, float))}
            break

    if len(rule.programmes) != 13 or len(rule.purposes) != 15:
        raise SystemExit(f"Expected 13 programmes and 15 purpose codes on {PROGRAM_SHEET}, found {len(rule.programmes)} and {len(rule.purposes)}")
    if not rule.published_total or not rule.gdp:
        raise SystemExit("The Program sheet's total row or the GDP sheet's nominal GDP row was not found")
    return rule


# ---------------------------------------------------------------------------
# The line sheets
# ---------------------------------------------------------------------------


@dataclass
class Item:
    """One counted line item, reduced to what the record needs."""

    year: str
    level: str
    line: str
    purpose: str
    plan: float
    actual: float
    counted: float  # the workbook's own after-BAR value
    rate: float  # the BAR rate the workbook applied (0 where its lookup found none)
    body: str = ""  # spending unit, in the sheet's own words
    body_en: str = ""
    holder: str = ""  # who holds the budget (Admin1, 2016-2021 sheet), its English label
    holder_mn: str = ""
    payment_en: str = ""  # economic classification, third level (2016-2021 sheet)
    payment_mn: str = ""


@dataclass
class Extract:
    items: list[Item] = field(default_factory=list)
    public: dict[tuple[str, str], float] = field(default_factory=dict)  # (year, level) -> MNT
    purpose_names_en: dict[str, str] = field(default_factory=dict)  # from Func2_EN, 2016-2021
    programme_names_en: dict[str, str] = field(default_factory=dict)  # from Prog2_EN, 2016-2021
    formula_2023: dict[str, float] = field(default_factory=dict)  # line -> what the 2023 formula sums, MNT
    dashes: int = 0  # actual cells holding "-" (SUMIFS reads them as nothing)


def _rate(value: Any) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


def read_2016_2021(wb: Workbook, rule: Rule, out: Extract) -> None:
    label_to_programme = {label.casefold(): code for code, label in rule.programme_labels.items()}
    purposes = set(rule.purposes)
    levels = {"01": "central", "02": "local"}
    for row, v in wb.rows("2016-2021"):
        if row < 3:
            continue
        year = excel_text(v.get("A"))
        if year not in CHECKED_YEARS:
            continue  # 2015 is not part of the review's totals
        level = levels.get(excel_text(v.get("B"))[:2])
        actual = v.get("Z")
        if not isinstance(actual, (int, float)) and actual is not None:
            out.dashes += 1
        if level:
            out.public[(year, level)] = out.public.get((year, level), 0.0) + number(actual)
        programme = label_to_programme.get(excel_text(v.get("T")).casefold())
        purpose = excel_text(v.get("X"))
        if programme:
            line = programme
        elif purpose in purposes and not_714(v.get("S")):
            line = purpose
        else:
            continue
        if not level:
            raise SystemExit(f"2016-2021 row {row}: a counted line item outside the State and local budgets ({v.get('B')!r})")
        out.items.append(
            Item(
                year=year,
                level=level,
                line=line,
                purpose=purpose,
                plan=number(v.get("Y")),
                actual=number(actual),
                counted=number(v.get("AB")),
                rate=_rate(v.get("AA")),
                body=excel_text(v.get("I")),
                body_en=excel_text(v.get("H")),
                holder=excel_text(v.get("D")),
                holder_mn=excel_text(v.get("E")),
                payment_en=excel_text(v.get("N")),
                payment_mn=excel_text(v.get("O")),
            )
        )
        if year in COMPLETE_YEARS:
            if purpose and v.get("AF"):
                out.purpose_names_en.setdefault(purpose, re.sub(r"^\d{5}\s+", "", excel_text(v.get("AF"))))
            if programme and v.get("AD"):
                out.programme_names_en.setdefault(programme, re.sub(r"^\d{5}\s+", "", excel_text(v.get("AD"))))


def read_2022(wb: Workbook, rule: Rule, out: Extract) -> None:
    programmes, purposes = set(rule.programmes), set(rule.purposes)
    for row, v in wb.rows("2022"):
        if row < 2:
            continue
        if code_of(v.get("A")) == "00211":  # total expenditure and net lending, one per report block
            out.public[("2022", "local")] = out.public.get(("2022", "local"), 0.0) + number(v.get("K"))
        programme = excel_text(v.get("T"))
        purpose = excel_text(v.get("V"))
        if programme in programmes:
            line = programme
        elif purpose in purposes and not_714(v.get("T")):
            line = purpose
        else:
            continue
        account = excel_text(v.get("R"))
        out.items.append(
            Item(
                year="2022",
                level="local",
                line=line,
                purpose=purpose,
                plan=number(v.get("F")),
                actual=number(v.get("G")),
                counted=number(v.get("X")),
                rate=_rate(v.get("W")),
                body=re.sub(r"^\d{6,}\s+", "", account),
            )
        )


def read_2023(wb: Workbook, rule: Rule, out: Extract) -> None:
    """A report by budget holder > programme > purpose, with no line items.

    The workbook's formula for a programme sums column I over the rows whose
    column B holds the programme's code (its programme rows); for a purpose
    code, over the rows whose B holds the code and whose A flag is not "714".
    The record reads the purpose rows, which carry the BAR rate, under their
    programme; `out.formula_2023` keeps what the formula itself sums, so the
    two readings are checked against each other."""
    programmes, purposes = set(rule.programmes), set(rule.purposes)
    holder, programme = "", ""
    for row, v in wb.rows("2023"):
        if row == 3:
            out.public[("2023", "central")] = number(v.get("E"))
            continue
        if row < 4:
            continue
        label = excel_text(v.get("C"))
        head = code_of(label)
        b = excel_text(v.get("B"))
        if re.fullmatch(r"\d{1,2}", head):
            holder, programme = re.sub(r"\s+", " ", label), ""
            continue
        if re.fullmatch(r"7\d{4}0", head):
            programme = head[:5]
            if b in programmes:
                out.formula_2023[b] = out.formula_2023.get(b, 0.0) + number(v.get("I"))
            continue
        if b in purposes and not_714(v.get("A")):
            out.formula_2023[b] = out.formula_2023.get(b, 0.0) + number(v.get("I"))
        if programme in programmes:
            line = programme
        elif b in purposes and not_714(v.get("A")):
            line = b
        else:
            continue
        out.items.append(
            Item(
                year="2023",
                level="central",
                line=line,
                purpose=b,
                plan=number(v.get("D")),
                actual=number(v.get("E")),
                counted=number(v.get("I")),
                rate=_rate(v.get("H")),
                body=holder,
            )
        )


def read_2024(wb: Workbook, rule: Rule, out: Extract) -> None:
    programmes, purposes = {int(c) for c in rule.programmes}, {int(c) for c in rule.purposes}
    national_total = None
    for row, v in wb.rows("2024"):
        if row < 2:
            continue
        if national_total is None and v.get("A") == 0 and code_of(v.get("G")) == "0211":
            national_total = number(v.get("O"))  # the national block's total expenditure and net lending
        y = v.get("Y") if isinstance(v.get("Y"), int) else None
        z = v.get("Z") if isinstance(v.get("Z"), int) else None
        if not z:
            # A report above purpose level (entity, or entity x programme):
            # its BAR lookup finds nothing, so it counts zero, and its plan
            # and actual repeat the purpose blocks below it.
            continue
        if y in programmes:
            line = str(y)
        elif z in purposes and not (y is not None and 71400 <= y <= 71499):
            line = str(z)
        else:
            continue
        out.items.append(
            Item(
                year="2024",
                level="local",
                line=line,
                purpose=str(z or ""),
                plan=number(v.get("K")),
                actual=number(v.get("L")),
                counted=number(v.get("AF")),
                rate=_rate(v.get("AE")),
                body=excel_text(v.get("D")),
            )
        )
    if national_total is None:
        raise SystemExit("2024: the national block's total expenditure row (0211) was not found")
    out.public[("2024", "local")] = national_total


def extract(path: Path, rule: Rule, wb: Workbook) -> Extract:
    stat = path.stat()
    CACHE.mkdir(parents=True, exist_ok=True)
    cached = CACHE / f"mongolia-extract-{stat.st_size}-{int(stat.st_mtime)}-v{RECORD_VERSION}-x{EXTRACT_VERSION}.pkl"
    if cached.exists():
        print(f"Line items from the cached extract {cached.name}")
        return pickle.loads(cached.read_bytes())
    out = Extract()
    for name, reader in (("2016-2021", read_2016_2021), ("2022", read_2022), ("2023", read_2023), ("2024", read_2024)):
        before = len(out.items)
        reader(wb, rule, out)
        print(f"  sheet {name}: {len(out.items) - before:,} counted line items")
    cached.write_bytes(pickle.dumps(out))
    return out


# ---------------------------------------------------------------------------
# The report: names and descriptions in the review's own words
# ---------------------------------------------------------------------------


def report_tables(path: Path) -> tuple[dict[str, tuple[str, str]], dict[str, tuple[str, str]], dict[str, str], dict[str, list[list[str]]]]:
    """(programmes, purposes outside, purpose names inside, captioned tables)
    from the report: Table 1 (code, name, description of the 13 programmes),
    Table 2 (the 15 purpose codes outside them), Tables 3-7 (purpose codes by
    attribution rate), and the rest by caption for the checks."""
    import docx  # python-docx, a pipeline dependency

    document = docx.Document(str(path))
    captions = [p.text.strip() for p in document.paragraphs if re.match(r"Table \d+\.", p.text.strip())]

    def cells(table) -> list[list[str]]:
        rows = []
        for r in table.rows:
            row: list[str] = []
            seen: set[int] = set()
            for c in r.cells:
                if id(c._tc) in seen:  # a merged cell is listed once per column it spans
                    continue
                seen.add(id(c._tc))
                row.append(" ".join(c.text.split()))
            rows.append(row)
        return rows

    tables = {caption: cells(t) for caption, t in zip(captions, document.tables)}
    programmes: dict[str, tuple[str, str]] = {}
    purposes: dict[str, tuple[str, str]] = {}
    inside: dict[str, str] = {}
    for caption, rows in tables.items():
        if caption.startswith("Table 1."):
            for row in rows:
                if len(row) >= 3 and re.fullmatch(r"714\d\d", row[0]):
                    programmes[row[0]] = (row[1], row[2])
        elif caption.startswith("Table 2."):
            for row in rows:
                if len(row) >= 3 and re.fullmatch(r"\d{5}", row[0]) and row[0] not in purposes:
                    purposes[row[0]] = (row[1], row[2])
        elif re.match(r"Table [3-7]\.", caption):
            for row in rows:
                if len(row) >= 2 and re.fullmatch(r"\d{5}", row[0]):
                    inside.setdefault(row[0], row[1])
    return programmes, purposes, inside, tables


def published_table(rows: list[list[str]], codes: list[str]) -> dict[tuple[str, str], tuple[float, float]]:
    """A report table of amounts by code or total row (rows) and year
    (columns): (printed value, half its last printed digit). A dash is zero.
    A row that stops early (the report merges its trailing dashes into one
    cell) leaves the remaining years unread."""
    header = next((r for r in rows if sum(1 for c in r if re.fullmatch(r"20\d\d", c)) >= 3), None)
    if header is None:
        return {}
    years = {i: c for i, c in enumerate(header) if re.fullmatch(r"20\d\d", c)}
    out: dict[tuple[str, str], tuple[float, float]] = {}
    for row in rows:
        label = row[0] if row else ""
        key = next((c for c in codes if label.startswith(c)), None)
        if key is None:
            low = label.lower()
            key = "total" if low.startswith("total") else "environment" if low == "environment" else "non-environmental" if low == "non-environmental" else None
        if key is None or len(row) != len(header):
            continue
        for i, year in years.items():
            printed = row[i].replace(",", "").replace(" ", "")
            if re.fullmatch(r"-?\d+(\.\d+)?", printed):
                decimals = len(printed.split(".")[1]) if "." in printed else 0
                out[(key, year)] = (float(printed), 0.5 * 10**-decimals + 1e-9)
            elif printed in ("-", "–"):
                out[(key, year)] = (0.0, 0.05)
    return out


# ---------------------------------------------------------------------------
# The record
# ---------------------------------------------------------------------------


def build(workbook: Path, report: Path) -> dict[str, Any]:
    with Workbook(workbook) as wb:
        rule = read_rule(wb)
        print(f"Rule read from {PROGRAM_SHEET}: {len(rule.programmes)} programmes, {len(rule.purposes)} purpose codes outside them")
        ex = extract(workbook, rule, wb)
    names_programme, names_purpose, names_inside, tables = report_tables(report)

    lines = rule.programmes + rule.purposes
    kind = {**{c: "programme" for c in rule.programmes}, **{c: "purpose" for c in rule.purposes}}

    # Sums the workbook's formulas take, line by line and year by year.
    counted: dict[tuple[str, str], float] = defaultdict(float)
    recomputed: dict[str, float] = defaultdict(float)
    cell: dict[tuple[str, str, str], dict[str, float]] = defaultdict(lambda: {"plan": 0.0, "actual": 0.0, "counted": 0.0})
    bodies: dict[tuple[str, str], dict[str, dict[str, Any]]] = defaultdict(dict)
    holders: dict[tuple[str, str], dict[str, dict[str, Any]]] = defaultdict(dict)
    holder_names: dict[str, set[tuple[str, str]]] = defaultdict(set)
    relevance: dict[str, dict[str, dict[str, float]]] = defaultdict(dict)
    payments: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
    for it in ex.items:
        counted[(it.line, it.year)] += it.counted
        recomputed[it.year] += it.rate * it.actual
        if it.year not in SHOWN_YEARS:
            continue
        c = cell[(it.line, it.year, it.level)]
        c["plan"] += it.plan
        c["actual"] += it.actual
        c["counted"] += it.counted
        if it.body:
            b = bodies[(it.line, it.year)].setdefault(it.body, {"name": it.body_en or "", "nameMn": it.body, "level": it.level, "counted": 0.0, "actual": 0.0})
            b["counted"] += it.counted
            b["actual"] += it.actual
        if it.year in COMPLETE_YEARS:
            # Who holds the budget, by the code that opens its label ("013 ...").
            code = code_of(it.holder)
            if not code or code != code_of(it.holder_mn):
                raise SystemExit(f"{it.year}, line {it.line}: a counted line item without one budget holder ({it.holder!r}, {it.holder_mn!r})")
            name, name_mn = (re.sub(r"^\d+\s+", "", label).strip() for label in (it.holder, it.holder_mn))
            holder_names[code].add((name, name_mn))
            h = holders[(it.line, it.year)].setdefault(code, {"code": code, "name": name, "nameMn": name_mn, "counted": 0.0})
            h["counted"] += it.counted
            r = relevance[it.line].setdefault(it.purpose, {"rate": it.rate, "actual": 0.0, "counted": 0.0})
            r["actual"] += it.actual
            r["counted"] += it.counted
            if it.payment_en:
                p = payments[it.line].setdefault(it.payment_en, {"nameMn": it.payment_mn, "counted": 0.0})
                p["counted"] += it.counted

    checks: list[Check] = []
    workbook_source = f"{WORKBOOK}, sheet {PROGRAM_SHEET}"
    for year in CHECKED_YEARS:
        total = sum(counted[(line, year)] for line in lines) / BN
        checks.append(Check(f"{year}, all lines", f"{workbook_source}, row 29", rule.published_total[year], total, 1e-6))
        checks.append(
            Check(
                f"{year}, counted = BAR rate x actual",
                f"{WORKBOOK}, the after-BAR column of the line sheets",
                total,
                recomputed[year] / BN,
                1e-6,
                rule="every year: counted = BAR rate x actual, line item by line item",
            )
        )
        for line in lines:
            checks.append(
                Check(f"{year}, {line}", workbook_source, rule.published[(line, year)], counted[(line, year)] / BN, 1e-6)
            )
    # 2023 is read from its purpose rows; the formula sums the programme rows.
    for line in lines:
        checks.append(
            Check(
                f"2023, {line}: purpose rows = programme row",
                f"{WORKBOOK}, sheet 2023",
                ex.formula_2023.get(line, 0.0) / BN,
                counted[(line, "2023")] / BN,
                1e-6,
                rule="2023: every line's purpose rows sum to the rows its formula reads",
            )
        )
    # The report prints the same figures, rounded (Tables 9 and 10). Its year
    # totals must match; a printed line figure that differs from the workbook
    # by more than its rounding is listed, not hidden.
    report_differences = []
    for prefix, codes in (("Table 9.", rule.programmes), ("Table 10.", rule.purposes)):
        caption = next((c for c in tables if c.startswith(prefix)), None)
        if caption is None:
            raise SystemExit(f"The report has no {prefix} table")
        printed = published_table(tables[caption], codes)
        sums = {
            "environment": lambda y: sum(counted[(c, y)] for c in rule.programmes),
            "non-environmental": lambda y: sum(counted[(c, y)] for c in rule.purposes),
            "total": (lambda y: sum(counted[(c, y)] for c in lines)) if prefix == "Table 9." else (lambda y: sum(counted[(c, y)] for c in rule.purposes)),
        }
        for (key, year), (value, tolerance) in sorted(printed.items()):
            reproduced = (sums[key](year) if key in sums else counted[(key, year)]) / BN
            if key == "total" and prefix == "Table 9.":
                checks.append(Check(f"{year}, all lines (report)", f"{REPORT}, {caption}, total row", value, reproduced, tolerance))
            elif abs(reproduced - value) > tolerance:
                report_differences.append(
                    {"what": f"{year}, {key}", "table": caption, "printed": value, "workbook": amount(reproduced)}
                )
    require_reproduced(checks)

    # The budget holders: one name per code, and in every line and year their
    # amounts add up to the line's counted amount, to the tögrög.
    for code, names in sorted(holder_names.items()):
        if len(names) > 1:
            raise SystemExit(f"Budget holder {code} has more than one name in {', '.join(COMPLETE_YEARS)}: {sorted(names)}")
    for line in lines:
        for year in COMPLETE_YEARS:
            held = sum(h["counted"] for h in holders[(line, year)].values())
            if abs(held - counted[(line, year)]) > 0.5:
                raise SystemExit(f"{year}, line {line}: its budget holders add up to MNT {held:,.2f}, not its counted MNT {counted[(line, year)]:,.2f}")

    rates = usd_rates("MNG")
    years_out = []
    for year in SHOWN_YEARS:
        in_data = SHEET_LEVELS[year]
        reviewed = {lvl: amount(sum(cell[(line, year, lvl)]["counted"] for line in lines) / BN) for lvl in in_data}
        public = {lvl: amount(ex.public[(year, lvl)] / BN) for lvl in in_data}
        years_out.append(
            {
                "year": year,
                "levels": in_data,
                "complete": in_data == ["central", "local"],
                "reviewed": reviewed,
                "public": public,
                "gdp": amount(rule.gdp[year]),
                "usdRate": rates[year],
            }
        )

    lines_out = []
    for line in lines:
        if kind[line] == "programme":
            name, description = names_programme[line]
            source = f"{REPORT}, Table 1"
            name_mn = re.sub(r"^\s*\d{5}\s+", "", rule.programme_labels[line]).strip()
        else:
            name, description = names_purpose[line]
            source = f"{REPORT}, Table 2"
            name_mn = rule.purpose_names_mn.get(line, "")
        amounts: dict[str, dict[str, Any]] = {}
        for year in SHOWN_YEARS:
            amounts[year] = {
                lvl: (
                    {k: amount(v / BN) for k, v in cell[(line, year, lvl)].items()}
                    if lvl in SHEET_LEVELS[year]
                    else None
                )
                for lvl, _ in LEVELS
            }
        body_years = {}
        for year in SHOWN_YEARS:
            found = sorted(bodies[(line, year)].values(), key=lambda b: (-b["counted"], b["nameMn"]))
            found = [b for b in found if abs(b["counted"]) > 0 or abs(b["actual"]) > 0]
            if found:
                body_years[year] = {
                    "count": len(found),
                    "top": [
                        {
                            **({"name": b["name"]} if b["name"] else {}),
                            "nameMn": b["nameMn"],
                            "level": b["level"],
                            "counted": amount(b["counted"] / BN),
                        }
                        for b in found[:BODIES_KEPT]
                    ],
                }
        holder_years = {}
        for year in COMPLETE_YEARS:
            found = sorted(holders[(line, year)].values(), key=lambda h: (-h["counted"], h["code"]))
            found = [h for h in found if abs(h["counted"]) > 0]
            if found:
                holder_years[year] = [{**h, "counted": amount(h["counted"] / BN)} for h in found]
        rel = []
        for purpose, r in sorted(relevance[line].items(), key=lambda kv: -kv[1]["actual"]):
            if abs(r["actual"]) < 1 and abs(r["counted"]) < 1:
                continue
            rel.append(
                {
                    "code": purpose,
                    "name": names_inside.get(purpose) or names_purpose.get(purpose, ("", ""))[0] or ex.purpose_names_en.get(purpose, ""),
                    "nameMn": rule.purpose_names_mn.get(purpose, ""),
                    "rate": r["rate"],
                    "actual": amount(r["actual"] / BN),
                    "counted": amount(r["counted"] / BN),
                }
            )
        paid = sum(p["counted"] for p in payments[line].values())
        pays = [
            {"name": re.sub(r"^\d{4}\s+", "", name_en), "nameMn": re.sub(r"^\d{4}\s+", "", p["nameMn"]), "share": round(p["counted"] / paid, 4)}
            for name_en, p in sorted(payments[line].items(), key=lambda kv: -kv[1]["counted"])[:PAYMENTS_KEPT]
            if paid > 0 and p["counted"] > 0
        ]
        lines_out.append(
            {
                "code": line,
                "kind": kind[line],
                "name": name,
                "nameMn": name_mn,
                **({"descriptionMn": rule.descriptions_mn[line]} if line in rule.descriptions_mn else {}),
                "description": description,
                "source": source,
                "amounts": amounts,
                "relevance": rel,
                "payments": pays,
                "bodies": body_years,
                "holders": holder_years,
            }
        )

    programs, expenditure = berdata_view(
        lines_out,
        COMPLETE_YEARS,
        type_of={**{c: "environmental" for c in rule.programmes}, **{c: "non_environmental" for c in rule.purposes}},
    )
    total_2024_bar_sheet = 196.41
    record = {
        "version": RECORD_VERSION,
        "review": {
            "name": "Biodiversity Expenditure Review (BER)",
            "title": "National Biodiversity Expenditure Review: Analysis of Public and Private Sector Biodiversity Expenditure (Progress Report II)",
            "preparedBy": "Unet Consulting LLC for UNDP's Biodiversity Finance Initiative (BIOFIN)",
            "date": "2025-11-10",
            "files": [WORKBOOK, REPORT],
            "workbookYears": [CHECKED_YEARS[0], CHECKED_YEARS[-1]],
        },
        "currency": "MNT",
        "unit": "billion",
        "usd": {"source": "World Bank, official exchange rate, annual average (PA.NUS.FCRF); python/data/fx-rates.json"},
        "levels": [{"id": lvl, "name": name} for lvl, name in LEVELS],
        "years": years_out,
        "lines": lines_out,
        "checks": checks_json(checks),
        "reportDifferences": report_differences,
        "notes": [
            *(
                [
                    f"The report's Tables 9 and 10 print {len(report_differences)} figures that differ from its workbook by more than rounding "
                    f"(at most MNT {max(abs(d['printed'] - d['workbook']) for d in report_differences):.2f} billion); the record follows the workbook, "
                    "whose totals the report's total row matches."
                ]
                if report_differences
                else []
            ),
            "The review's workbook holds different parts of the budget each year: 2020 and 2021 the State and the local budgets, 2022 and 2024 the local budgets only, 2023 the State budget only. A year total is given for 2020 and 2021 only.",
            "Counted is the review's own rule: every line item of the 13 environment programmes at the attribution rate (BAR) of its purpose code, plus 15 environmental purpose codes in other programmes at 100%. Core operating costs count at 100%; capital works coded 80918 at 5%.",
            "All public spending is the State budget and the local budgets, each as executed (gross); transfers from the State budget to local budgets are not netted out.",
            f"2024: the workbook's BAR sheet and the report's Table 11 give MNT {total_2024_bar_sheet} billion, MNT 19.73 billion less than its Program sheet and Table 9, because a helper column those totals filter on is blank in the 2024 sheet's first 130,690 rows. The record follows the Program sheet.",
            "Line names and descriptions are the report's own (Tables 1 and 2); spending units are named as in the workbook, in English where its 2016-2021 sheet gives English.",
            "Budget holders (ministers, the capital city, the provinces) are the 2016-2021 sheet's own (column Admin1, in its English and Mongolian, the leading code kept apart), for 2020 and 2021; in each year their amounts add up to the line's counted amount.",
        ],
    }
    return {
        "_provenance": (
            f"Built by python/scripts/build_ber_mongolia.py from the review's workbook ({WORKBOOK}) and report ({REPORT}). "
            "Amounts follow the review's own rule, read from its formulas; every published total is reproduced (record.checks). "
            "No AI touched any figure, name or description."
        ),
        "programs": programs,
        "expenditure": expenditure,
        "currency": "MNT",
        "unit": "billion",
        "period": {"start": int(COMPLETE_YEARS[0]), "end": int(COMPLETE_YEARS[-1])},
        "record": record,
    }


def sharepoint_file(name: str) -> Path:
    roots = glob.glob(os.path.expanduser("~/Library/CloudStorage/*/*AIFlagship*"))
    if not roots:
        raise SystemExit("The SharePoint folder is not mounted; pass --workbook and --report")
    return Path(roots[0]) / FOLDER / name


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, help=f"{WORKBOOK} (default: the SharePoint copy)")
    parser.add_argument("--report", type=Path, help=f"{REPORT} (default: the SharePoint copy)")
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--check-only", action="store_true", help="Reconcile and report; write nothing")
    args = parser.parse_args()
    workbook = args.workbook or sharepoint_file(WORKBOOK)
    report = args.report or sharepoint_file(REPORT)
    payload = build(workbook, report)
    record = payload["record"]
    print(reconciliation_markdown("Mongolia", record).split("\n## Notes")[0][:1200])
    ok = sum(1 for c in record["checks"] if c["ok"])
    print(f"{ok} of {len(record['checks'])} published totals reproduced")
    if not args.check_only:
        write_record(args.out, payload)
        (CACHE / "mongolia-reconciliation.md").write_text(reconciliation_markdown("Mongolia", record))
        print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
