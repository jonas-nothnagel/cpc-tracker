"""BIOFIN's GLOBE expenditure types, with the GBF targets the workbook files
each under, for the official chain of the BER plan (step 3): a budget line's
expenditure type leads to GBF targets, and the national targets a country
filed under the same GBF targets in its 7th National Report.

Reads the GLOBE workbook (BIOFIN, 2024; dev_data_scripts/biofin/
2-8_taxonomy_globe.xlsx, not in git) and writes python/data/
globe-expenditure-types.json: every level-3 expenditure programme of the nine
category sheets (and its lettered sub-items), verbatim:

  - name: column C, "Expenditure programs";
  - definition: column D, "Definitions / Examples";
  - examples: column E, the expenditure it may include;
  - gbf: the GBF targets column M lists ("1 - Spatial planning", one per
    line), as numbers; the target names are the CBD's own, kept in
    gbfTargets from the workbook's Matching sheet.

Nothing is drafted or inferred: a type without a GBF entry keeps an empty list.

Run from python/:
    .venv/bin/python scripts/build_globe_types.py [--workbook PATH]
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "dev_data_scripts" / "biofin" / "2-8_taxonomy_globe.xlsx"
OUT = ROOT / "python" / "data" / "globe-expenditure-types.json"

TYPE_CODE = re.compile(r"^(\d)\.(\d\d)\.(\d\d)([a-z])?\)?$")
GBF_ENTRY = re.compile(r"^\s*(\d{1,2})\s*[-–]\s*\S")

TERMS = (
    "BIOFIN, The Biodiversity Finance Initiative: GLOBE taxonomy workbook (2024). "
    "Copyright © 2024. All rights reserved. Reproduction for educational or other "
    "non-commercial purposes is authorized without prior written permission provided "
    "the source is fully acknowledged. Whether the open-source tool may carry this "
    "table is open with BIOFIN (BER handover, 23 October 2026)."
)


def text(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def gbf_numbers(cell) -> list[int]:
    """The GBF targets a GBF Alignment cell lists, one "N - Title" per line."""
    found = []
    for line in str(cell or "").splitlines():
        m = GBF_ENTRY.match(line)
        if m and 1 <= int(m.group(1)) <= 23:
            found.append(int(m.group(1)))
    return sorted(set(found))


def build(workbook: Path) -> dict:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    types = []
    for ws in wb.worksheets:
        if not re.match(r"^\d ", ws.title):
            continue
        rows = list(ws.iter_rows(values_only=True))
        header = [text(h).lower() for h in rows[1]]
        gbf_col = next(i for i, h in enumerate(header) if h.startswith("gbf"))
        assert header[2].startswith("expenditure programs") and header[3].startswith("definitions"), (ws.title, header)
        for r in rows[2:]:
            m = TYPE_CODE.match(text(r[0]))
            if not m:
                continue
            types.append({
                "id": f"{m.group(1)}.{m.group(2)}.{m.group(3)}{m.group(4) or ''}",
                "category": m.group(1),
                "subcategory": f"{m.group(1)}.{m.group(2)}",
                "sheet": ws.title,
                "name": text(r[2]),
                "definition": text(r[3]),
                "examples": text(r[4]),
                "gbf": gbf_numbers(r[gbf_col]),
            })
    gbf_targets: dict[str, str] = {}
    for r in list(wb["Matching"].iter_rows(values_only=True))[1:]:
        m = re.match(r"^\s*(\d{1,2})\s*[-–]\s*(.+)$", text(r[1]))
        if m and m.group(1) not in gbf_targets:
            gbf_targets[m.group(1)] = m.group(2).strip()
    ids = [t["id"] for t in types]
    assert len(ids) == len(set(ids)), "duplicate type codes"
    return {
        "_source": (
            "BIOFIN GLOBE 2024 taxonomy workbook (dev_data_scripts/biofin/2-8_taxonomy_globe.xlsx), "
            "the nine category sheets: every level-3 expenditure programme and lettered sub-item, "
            "name (column C), definition (D), examples (E) and GBF alignment (M) verbatim; GBF "
            "target names from its Matching sheet. Built by python/scripts/build_globe_types.py."
        ),
        "_terms": TERMS,
        "gbfTargets": dict(sorted(gbf_targets.items(), key=lambda kv: int(kv[0]))),
        "types": types,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    args = parser.parse_args()
    data = build(args.workbook)
    OUT.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n")
    types = data["types"]
    print(f"{len(types)} expenditure types, {sum(1 for t in types if t['gbf'])} with GBF targets -> {OUT}")


if __name__ == "__main__":
    main()
