"""Build Sri Lanka's spending record from its Biodiversity Expenditure Review (BER).

Step 1 of the BER plan, "Show the money". Reads the review's dataset (BER
2015-2024 Dataset .xlsx, the source of record: the tidier copy shared on 15
October 2026 drops institution and project names, gross amounts and weights)
and writes python/data/sri-lanka-ber.json. No AI touches a figure, a name or a
description.

The review counts in two ways, and the record keeps them apart, never added:

  - institutions (the headline): the four core agencies and the dedicated
    biodiversity units in full, the other agencies at their stated relevance
    rate. Their sum is the review's total, sheet All Central Govt.
  - project lines: each budget line of the review's project sheet at its own
    weight (gross x weightage), with the review's tags and tier. Many sit
    inside the institutions above, so they are a second view of the same
    money, and the lines the AI places in policy areas (step 2).

Every year's total and every agency group's must equal sheet All Central Govt,
every relevant amount must equal gross x relevance, and every weighted project
amount gross x weight, or nothing is written.

Run from python/:
    .venv/bin/python scripts/build_ber_sri_lanka.py
"""

from __future__ import annotations

import argparse
import glob
import os
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.ber_record import (  # noqa: E402
    RECORD_VERSION,
    Check,
    amount,
    berdata_view,
    checks_json,
    reconciliation_markdown,
    require_reproduced,
    write_record,
)
from src.config import DATA_DIR  # noqa: E402

FOLDER = "Sri Lanka/Expenditure data"
DATASET = "BER 2015-2024 Dataset .xlsx"
OUT = DATA_DIR / "sri-lanka-ber.json"
CACHE = Path(__file__).resolve().parent.parent / "tmp" / "ber"

YEARS = [str(y) for y in range(2015, 2025)]
PROJECT_YEARS = [str(y) for y in range(2014, 2025)]  # the project sheet starts a year earlier
EXACT = 1e-6  # LKR million: a rounding allowance only

GROUPS = [
    ("CBA", "Core biodiversity agencies"),
    ("DBU", "Dedicated biodiversity units"),
    ("OASBE", "Other agencies with significant biodiversity expenditure"),
    ("OAMBE", "Other agencies with marginal biodiversity expenditure"),
]
GROUP_ROWS = {"CBA": 2, "DBU": 3, "OASBE": 4, "OAMBE": 5, "total": 6}


def split_name(label: str) -> tuple[str, str]:
    """"Forest Department (FD)" -> ("Forest Department", "FD")."""
    label = " ".join(str(label).split())
    match = re.match(r"^(.*?)\s*\(([A-Z][A-Za-z-]*)\)\s*(.*)$", label)
    if match:
        name = " ".join(f"{match.group(1)} {match.group(3)}".split())
        return name, match.group(2)
    return label, ""


def number(value: Any) -> float:
    return float(value) if isinstance(value, (int, float)) else 0.0


def build(path: Path) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    checks: list[Check] = []

    # --- The headline sheet: totals, GDP, exchange rate, all government spending
    head = wb["All Central Govt"]
    cols = {str(head.cell(1, c).value): c for c in range(2, 12) if isinstance(head.cell(1, c).value, int)}
    if list(cols) != YEARS:
        raise SystemExit(f"All Central Govt: expected the years {YEARS} in row 1, found {list(cols)}")

    def row_values(label_start: str, row: int) -> dict[str, float]:
        label = str(head.cell(row, 1).value or "")
        if not label.startswith(label_start):
            raise SystemExit(f"All Central Govt row {row}: expected {label_start!r}, found {label!r}")
        return {y: number(head.cell(row, c).value) for y, c in cols.items()}

    published = {
        "CBA": row_values("Core Biodiversity Agencies", 2),
        "DBU": row_values("Dedicated Biodiversity Units", 3),
        "OASBE": row_values("Other Agencies with Significant", 4),
        "OAMBE": row_values("Other Agencies with Marginal", 5),
        "total": row_values("Total", 6),
    }
    gdp = row_values("NGDP", 14)
    usd = row_values("USD/LKR", 23)
    government = row_values("Total Primary Expenditure, LKR", 33)

    # --- Institutions, by agency group
    institutions: dict[str, dict[str, Any]] = {}

    def institution(group: str, label: str, sheet: str) -> dict[str, Any]:
        name, acronym = split_name(label)
        key = acronym or name
        entry = institutions.setdefault(
            key,
            {"code": key, "kind": "institution", "group": group, "name": name, "acronym": acronym, "source": f"{DATASET}, sheet {sheet}", "years": {}},
        )
        return entry

    ws = wb["CBA Exp"]
    for r in range(3, ws.max_row + 1):
        year, label = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not isinstance(year, int) or not label or str(label).startswith("Total"):
            continue
        inst = institution("CBA", label, "CBA Exp")
        gross = number(ws.cell(r, 8).value)
        inst["years"][str(year)] = {
            "plan": number(ws.cell(r, 5).value),
            "gross": gross,
            "rate": 1.0,
            "counted": gross,
            "recurrent": number(ws.cell(r, 6).value),
            "capital": number(ws.cell(r, 7).value),
        }

    ws = wb["DBU Exp"]  # in rupees, not millions
    for r in range(4, ws.max_row + 1):
        year, label = ws.cell(r, 1).value, ws.cell(r, 2).value
        if not isinstance(year, int) or not label or str(label).strip().startswith("Total"):
            continue
        inst = institution("DBU", label, "DBU Exp")
        gross = number(ws.cell(r, 10).value) / 1e6
        inst["years"][str(year)] = {
            "plan": number(ws.cell(r, 6).value) / 1e6,
            "gross": gross,
            "rate": 1.0,
            "counted": gross,
        }

    rationale: dict[str, str] = {}
    for sheet, group in (("OASigBE Exp", "OASBE"), ("OAMarBE Exp", "OAMBE")):
        ws = wb[sheet]
        for r in range(4, ws.max_row + 1):
            year, label = ws.cell(r, 1).value, ws.cell(r, 2).value
            if not isinstance(year, int) or not label:
                continue
            inst = institution(group, label, sheet)
            gross = number(ws.cell(r, 5).value)
            rate = number(ws.cell(r, 6).value)
            relevant = number(ws.cell(r, 7).value)
            checks.append(Check(f"{year}, {inst['code']}: relevant = gross x rate", f"{DATASET}, sheet {sheet}, row {r}", relevant, gross * rate, EXACT, rule="every relevant amount = gross x relevance rate"))
            inst["years"][str(year)] = {"gross": gross, "rate": rate, "counted": relevant}
            note = ws.cell(r, 8).value
            if isinstance(note, str) and note.strip() and inst["code"] not in rationale:
                rationale[inst["code"]] = " ".join(note.split())

    # The review's 2024 figure for the dedicated units is a single estimate on
    # its headline sheet, with no breakdown by unit.
    dbu_2024 = published["DBU"]["2024"]
    if dbu_2024 and not any("2024" in i["years"] for i in institutions.values() if i["group"] == "DBU"):
        institutions["DBU-2024"] = {
            "code": "DBU-2024",
            "kind": "institution",
            "group": "DBU",
            "name": "Dedicated biodiversity units, 2024 (one estimate, not broken down by unit)",
            "acronym": "",
            "source": f"{DATASET}, sheet All Central Govt, row 3",
            "years": {"2024": {"gross": dbu_2024, "rate": 1.0, "counted": dbu_2024}},
        }

    for key, label in GROUPS:
        for year in YEARS:
            got = sum(i["years"].get(year, {}).get("counted", 0.0) for i in institutions.values() if i["group"] == key)
            checks.append(Check(f"{year}, {key}", f"{DATASET}, sheet All Central Govt, row {GROUP_ROWS[key]}", published[key][year], got, EXACT))
    for year in YEARS:
        got = sum(i["years"].get(year, {}).get("counted", 0.0) for i in institutions.values())
        checks.append(Check(f"{year}, all institutions", f"{DATASET}, sheet All Central Govt, row 6", published["total"][year], got, EXACT))

    # --- Project lines
    ws = wb["Govt Project wise Tagging"]
    header = [ws.cell(2, c).value for c in range(1, 33)]
    expected = ["Institution", "Project", *map(int, PROJECT_YEARS), "2025 Allo", "2025 Act", "2015-2024", "Tag 1", "Tag 2", "Tag 3", "Weightage", "Expenditure Tag", *map(int, PROJECT_YEARS)]
    if header != expected:
        raise SystemExit(f"Govt Project wise Tagging: the header row has changed:\n  found    {header}\n  expected {expected}")
    projects = []
    for r in range(3, 109):
        title = ws.cell(r, 2).value
        if not title:
            continue
        gross = {y: number(ws.cell(r, 3 + i).value) for i, y in enumerate(PROJECT_YEARS)}
        weight = number(ws.cell(r, 20).value)
        weighted = {y: number(ws.cell(r, 22 + i).value) for i, y in enumerate(PROJECT_YEARS)}
        for y in PROJECT_YEARS:
            checks.append(Check(f"{y}, project row {r}: weighted = gross x weight", f"{DATASET}, sheet Govt Project wise Tagging, row {r}", weighted[y], gross[y] * weight, 1e-3, rule="every weighted project amount = gross x weight"))
        tags = [" ".join(str(ws.cell(r, c).value).split()) for c in (17, 18, 19) if ws.cell(r, c).value]
        projects.append(
            {
                "code": f"P{r:03d}",
                "kind": "project",
                "name": " ".join(str(title).split()),
                "ministry": " ".join(str(ws.cell(r, 1).value or "").split()),
                "source": f"{DATASET}, sheet Govt Project wise Tagging, row {r}",
                "tags": tags,
                "tier": " ".join(str(ws.cell(r, 21).value or "").split()),
                "weight": weight,
                "years": {y: {"gross": gross[y], "rate": weight, "counted": weighted[y]} for y in PROJECT_YEARS},
                "allocation2025": number(ws.cell(r, 14).value),
                "actual2025": number(ws.cell(r, 15).value),
            }
        )
    if len(projects) != 106:
        raise SystemExit(f"Expected 106 project lines, found {len(projects)}")
    require_reproduced(checks)

    # One pair found in the October 2026 assessment, not by a rule: the World
    # Bank's ecosystem conservation project (ESCAMP) under two titles. The
    # review counts both; so does the record, and the pair is flagged for the
    # review team to confirm.
    for p in projects:
        other = {"P045": "P071", "P071": "P045"}.get(p["code"])
        if other:
            p["sameProjectAs"] = other
    pair = {p["code"]: p for p in projects if p["code"] in ("P045", "P071")}
    if set(pair) != {"P045", "P071"} or any(
        abs(pair["P045"]["years"][y]["gross"] - pair["P071"]["years"][y]["gross"]) > 0.05 for y in ("2017", "2018", "2019")
    ):
        raise SystemExit("Rows 45 and 71 of the project sheet no longer hold the ESCAMP pair; review the flag")

    # --- The record
    def line_amounts(years: dict[str, dict[str, float]]) -> dict[str, Any]:
        return {
            y: {"all": {k: amount(v) for k, v in years[y].items() if k in ("plan", "gross", "rate", "counted")}}
            for y in YEARS
            if y in years
        }

    institution_lines = []
    for inst in institutions.values():
        institution_lines.append(
            {
                "code": inst["code"],
                "kind": "institution",
                "group": inst["group"],
                "name": inst["name"],
                **({"acronym": inst["acronym"]} if inst["acronym"] else {}),
                **({"rationale": rationale[inst["code"]]} if inst["code"] in rationale else {}),
                "source": inst["source"],
                "amounts": line_amounts(inst["years"]),
            }
        )
    project_lines = []
    for p in projects:
        project_lines.append(
            {
                "code": p["code"],
                "kind": "project",
                "name": p["name"],
                "description": "",
                "ministry": p["ministry"],
                "tags": p["tags"],
                "tier": p["tier"],
                "weight": p["weight"],
                **({"sameProjectAs": p["sameProjectAs"]} if p.get("sameProjectAs") else {}),
                "source": p["source"],
                "amounts": line_amounts(p["years"]),
                "allocation2025": amount(p["allocation2025"]),
                "actual2025": amount(p["actual2025"]),
            }
        )
    project_lines.sort(key=lambda l: (-sum(c["all"]["counted"] for c in l["amounts"].values()), l["code"]))

    years_out = []
    for year in YEARS:
        by_group = {
            key: amount(sum(i["years"].get(year, {}).get("counted", 0.0) for i in institutions.values() if i["group"] == key))
            for key, _ in GROUPS
        }
        years_out.append(
            {
                "year": year,
                "levels": ["all"],
                "complete": True,
                "reviewed": {"all": amount(published["total"][year])},
                "byGroup": by_group,
                "public": {"all": amount(government[year])},
                "gdp": amount(gdp[year]),
                "usdRate": usd[year],
                "projects": amount(sum(p["years"][year]["counted"] for p in projects)),
            }
        )

    programs, expenditure = berdata_view(project_lines, YEARS)
    for prog, line in zip(programs, project_lines):
        prog["institution"] = line["ministry"]
    record = {
        "version": RECORD_VERSION,
        "review": {
            "name": "Biodiversity Expenditure Review (BER)",
            "title": "Biodiversity Expenditure Review, Sri Lanka 2015-2024 (dataset)",
            "preparedBy": "The review team for UNDP's Biodiversity Finance Initiative (BIOFIN), Sri Lanka",
            "date": "2026-08-31",
            "files": [DATASET],
            "workbookYears": [YEARS[0], YEARS[-1]],
        },
        "currency": "LKR",
        "unit": "million",
        "usd": {"source": "The review's own annual average rates (sheet All Central Govt, row 23)."},
        "levels": [{"id": "all", "name": "Central government"}],
        "groups": [{"id": key, "name": label} for key, label in GROUPS],
        "years": years_out,
        "lines": institution_lines + project_lines,
        "checks": checks_json(checks),
        "notes": [
            "The review's total counts institutions: the four core biodiversity agencies and the dedicated units in full, other agencies at the relevance rate the review states for each.",
            "The review's project lines are a second view of the same money, each at its own weight; many sit inside the institutions above, so the two are never added. The policy areas are read from the project lines, where the review's tags sit.",
            "All government spending is the series the review labels total primary expenditure; it equals the Ministry of Finance's total expenditure, interest included (2023: LKR 5,357 billion).",
            "US$ figures use the review's own annual average rates; in US$ the series falls after the 2022 devaluation while it rises in rupees.",
            "The dedicated units' 2024 figure is a single estimate of LKR 120 million, not broken down by unit; Sri Lanka Customs has no figures for 2022-2024 and the Land Development Corporation none for 2024.",
            "Project rows 45 and 71 appear to be the same World Bank project (ecosystem conservation and management) under two titles: equal amounts in 2018 and 2019 and LKR 203.4 against 203.431 million in 2017, LKR 1,896.6 million in all. The review counts both, as does the record; the pair is flagged for the review team to confirm.",
        ],
    }
    return {
        "_provenance": (
            f"Built by python/scripts/build_ber_sri_lanka.py from the review's dataset ({DATASET}). "
            "Amounts follow the review's own rules; every total on its headline sheet is reproduced (record.checks). "
            "No AI touched any figure, name or description."
        ),
        "programs": programs,
        "expenditure": expenditure,
        "currency": "LKR",
        "unit": "million",
        "period": {"start": int(YEARS[0]), "end": int(YEARS[-1])},
        "record": record,
    }


def sharepoint_file(name: str) -> Path:
    roots = glob.glob(os.path.expanduser("~/Library/CloudStorage/*/*AIFlagship*"))
    if not roots:
        raise SystemExit("The SharePoint folder is not mounted; pass --dataset")
    return Path(roots[0]) / FOLDER / name


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--dataset", type=Path, help=f"{DATASET} (default: the SharePoint copy)")
    parser.add_argument("--out", type=Path, default=OUT)
    parser.add_argument("--check-only", action="store_true", help="Reconcile and report; write nothing")
    args = parser.parse_args()
    payload = build(args.dataset or sharepoint_file(DATASET))
    record = payload["record"]
    ok = sum(1 for c in record["checks"] if c["ok"])
    kinds = {k: sum(1 for l in record["lines"] if l["kind"] == k) for k in ("institution", "project")}
    print(f"{ok} of {len(record['checks'])} published totals reproduced; {kinds['institution']} institutions, {kinds['project']} project lines")
    if not args.check_only:
        write_record(args.out, payload)
        CACHE.mkdir(parents=True, exist_ok=True)
        (CACHE / "sri-lanka-reconciliation.md").write_text(reconciliation_markdown("Sri Lanka", record))
        print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
