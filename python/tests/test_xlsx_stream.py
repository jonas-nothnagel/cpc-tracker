"""The streaming .xlsx reader reads what Excel stored: values, text and formulas."""

from __future__ import annotations

import openpyxl

from src.xlsx_stream import Workbook, column_letters


def make_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Хөтөлбөр"
    ws.append(["Code", "Name", 2020])
    ws.append(["71405", "Тусгай хамгаалалттай газар", 10.205471251304992])
    ws.append([81707, "Air pollution reduction", 2])
    ws["D2"] = "=C2*2"
    other = wb.create_sheet("2022")
    other.append(["A", "B"])
    other.append([None, "only B"])
    wb.save(path)


def test_rows_give_numbers_text_and_skip_empty_cells(tmp_path):
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    with Workbook(path) as wb:
        assert list(wb.sheets) == ["Хөтөлбөр", "2022"]
        rows = dict(wb.rows("Хөтөлбөр"))
        assert rows[1] == {"A": "Code", "B": "Name", "C": 2020}
        assert rows[2]["A"] == "71405" and abs(rows[2]["C"] - 10.205471251304992) < 1e-12
        assert rows[3]["A"] == 81707 and isinstance(rows[3]["A"], int)
        # openpyxl writes no cached value for a formula: the cell is left out.
        assert "D" not in rows[2]
        assert dict(wb.rows("2022"))[2] == {"B": "only B"}


def test_formulas_are_read_as_written(tmp_path):
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    with Workbook(path) as wb:
        assert dict(wb.formulas("Хөтөлбөр")) == {"D2": "C2*2"}


def test_max_rows_stops_early(tmp_path):
    path = tmp_path / "book.xlsx"
    make_workbook(path)
    with Workbook(path) as wb:
        assert [r for r, _ in wb.rows("Хөтөлбөр", max_rows=2)] == [1, 2]


def test_column_letters():
    assert column_letters("AB12") == "AB"
    assert column_letters("C3") == "C"
